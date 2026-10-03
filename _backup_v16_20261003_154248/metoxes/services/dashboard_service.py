from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]
DASHBOARD_JSON_FILE = BASE_DIR / "portfolio_dashboard_data.json"
DASHBOARD_HTML_FILE = BASE_DIR / "portfolio_dashboard.html"
PORTFOLIO_EXCEL_FILE = BASE_DIR / "portfolio.xlsx"
CANDIDATES_SHEET_NAME = "Candidates"

HIDDEN_SCORE_FIELDS = {
    "relative_score",
    "absolute_score",
    "absolute_coverage",
    "score_mode",
}


def _json_default(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()

    raise TypeError(
        f"Object of type {type(value).__name__} "
        "is not JSON serializable"
    )


def _public_stock_view(stock):
    return {
        key: value
        for key, value in stock.items()
        if key not in HIDDEN_SCORE_FIELDS
        and not key.startswith("absolute_")
        and not key.startswith("relative_")
        and "_absolute_score" not in key
        and "_relative_score" not in key
    }


def _serializable_stocks(stocks):
    public_rows = [
        _public_stock_view(stock)
        for stock in stocks or []
    ]

    return json.loads(
        json.dumps(
            public_rows,
            ensure_ascii=False,
            default=_json_default,
        )
    )


def _load_candidates_from_excel():
    if not PORTFOLIO_EXCEL_FILE.exists():
        return []

    try:
        workbook = load_workbook(
            PORTFOLIO_EXCEL_FILE,
            data_only=True,
            read_only=True,
        )
    except Exception:
        return []

    try:
        if CANDIDATES_SHEET_NAME not in workbook.sheetnames:
            return []

        sheet = workbook[CANDIDATES_SHEET_NAME]
        headers = {}

        for cell in sheet[1]:
            if cell.value is None:
                continue

            headers[
                str(cell.value).strip().lower()
            ] = cell.column

        wanted = (
            "rank",
            "symbol",
            "name",
            "sector",
            "country",
            "current_price",
            "high_52w",
            "price_25pct_below_high",
            "discount_from_52w_high_pct",
            "final_score",
            "validation_status",
            "dashboard_analysis",
        )

        rows = []

        for row_number in range(2, sheet.max_row + 1):
            symbol_col = headers.get("symbol")

            if symbol_col is None:
                break

            symbol = sheet.cell(
                row=row_number,
                column=symbol_col,
            ).value

            if not symbol:
                continue

            row = {}

            for field in wanted:
                column = headers.get(field)

                if column is not None:
                    row[field] = sheet.cell(
                        row=row_number,
                        column=column,
                    ).value

            rows.append(row)

        return rows
    finally:
        workbook.close()


HTML_TEMPLATE = r'''<!DOCTYPE html>
<html lang="el">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>METOXES2 Portfolio Dashboard</title>
<script src="https://cdn.plot.ly/plotly-3.3.1.min.js"></script>
<style>
:root{--blue:#4472C4;--blue2:#2F5597;--bg:#F4F7FB;--card:#fff;--text:#172033;--muted:#697386;--border:#DCE3ED;--green:#16845B;--red:#C43D4B;--orange:#D97706}
*{box-sizing:border-box}body{margin:0;font-family:Segoe UI,Arial,sans-serif;background:var(--bg);color:var(--text)}
.header{background:linear-gradient(135deg,var(--blue2),var(--blue));color:#fff;padding:22px 28px}.header h1{margin:0;font-size:26px}.header p{margin:7px 0 0;opacity:.88}
.wrap{max-width:1550px;margin:auto;padding:18px}.filters{display:grid;grid-template-columns:repeat(8,minmax(135px,1fr));gap:10px;background:var(--card);padding:14px;border:1px solid var(--border);border-radius:12px;margin-bottom:14px}
label{font-size:12px;color:var(--muted);font-weight:600;display:block;margin-bottom:5px}select,input{width:100%;padding:8px;border:1px solid var(--border);border-radius:8px;background:#fff}
.kpis{display:grid;grid-template-columns:repeat(7,1fr);gap:10px;margin-bottom:14px}.kpi,.card{background:var(--card);border:1px solid var(--border);border-radius:12px}.kpi{padding:14px}.kpi .l{font-size:11px;color:var(--muted);text-transform:uppercase}.kpi .v{font-size:22px;font-weight:700;margin-top:7px}
.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}.card{padding:8px}.full{grid-column:1/-1}.title{font-weight:700;padding:10px 12px 0}.chart{height:410px}.tall{height:510px}
.tablewrap{max-height:560px;overflow:auto;margin:8px;border:1px solid var(--border);border-radius:8px}
table{border-collapse:collapse;width:100%;font-size:12px}th{position:sticky;top:0;background:var(--blue);color:#fff;padding:8px;white-space:nowrap}td{padding:8px;border-bottom:1px solid #E9EDF3;white-space:nowrap}.pos{color:var(--green);font-weight:600}.neg{color:var(--red);font-weight:600}
.scoretable table{min-width:0}.scoretable .tablewrap{max-height:520px}.detail table{min-width:1750px}
.analysis-list{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px;padding:10px}.analysis-item{border:1px solid var(--border);border-radius:10px;padding:14px;background:#FAFCFF}.analysis-item h3{font-size:14px;margin:0 0 7px}.analysis-meta{font-size:12px;color:var(--muted);margin-bottom:8px}.analysis-item p{font-size:13px;line-height:1.55;margin:0}
.accordion-list{display:grid;gap:9px;padding:10px}.stock-details{border:1px solid var(--border);border-radius:11px;background:#fff;overflow:hidden}.stock-details[open]{box-shadow:0 5px 18px rgba(31,78,121,.08)}.stock-summary{display:grid;grid-template-columns:48px 110px minmax(180px,1fr) 120px 125px;gap:10px;align-items:center;padding:12px 14px;cursor:pointer;list-style:none}.stock-summary::-webkit-details-marker{display:none}.stock-summary:hover{background:#F8FAFD}.stock-rank{color:var(--muted);font-size:12px}.stock-symbol{font-weight:800}.stock-name{font-weight:600}.score-badge{display:inline-flex;justify-content:center;align-items:center;border-radius:999px;padding:6px 10px;font-weight:800;background:#EAF1FB;color:var(--blue2)}.source-badge{font-size:11px;text-align:center;border:1px solid var(--border);border-radius:999px;padding:5px 8px;color:var(--muted)}.stock-panel{border-top:1px solid var(--border);padding:14px;background:#FBFCFE}.analysis-grid{display:grid;grid-template-columns:1.2fr 1fr;gap:12px}.subcard{border:1px solid #E5EAF1;border-radius:9px;background:#fff;padding:12px}.subcard h4{margin:0 0 8px;font-size:13px}.subcard p{font-size:12.5px;line-height:1.55;margin:4px 0}.metric-grid{display:grid;grid-template-columns:repeat(5,minmax(110px,1fr));gap:8px;margin-top:10px}.metric-chip{border:1px solid #E5EAF1;border-radius:8px;padding:8px;background:#fff}.metric-chip .m{font-size:11px;color:var(--muted)}.metric-chip .v{font-size:14px;font-weight:700;margin-top:3px}.impact-positive{color:var(--green)}.impact-negative{color:var(--red)}.impact-neutral{color:var(--orange)}.bullet-line{margin:4px 0;padding-left:16px}.bullet-line li{margin:3px 0;font-size:12.5px;line-height:1.45}.news-list{margin:5px 0 0;padding-left:16px}.news-list li{margin:3px 0;font-size:12.5px;line-height:1.4}.status-warning{color:var(--orange);font-weight:700}.status-error{color:var(--red);font-weight:700}.status-ok{color:var(--green);font-weight:700}
@media(max-width:1200px){.filters{grid-template-columns:repeat(4,1fr)}.kpis{grid-template-columns:repeat(4,1fr)}}@media(max-width:900px){.filters,.kpis,.grid,.analysis-list,.analysis-grid{grid-template-columns:1fr}.stock-summary{grid-template-columns:34px 80px 1fr 88px}.stock-summary .source-badge{display:none}.metric-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.full{grid-column:auto}}
</style>
</head>
<body>
<div class="header"><h1>METOXES2 Portfolio Dashboard</h1><p>Portfolio • υποψήφιες μετοχές • συνολικό score • 52-week high • data quality</p></div>
<div class="wrap">

<div class="filters">
<div><label>Sector</label><select id="sector"></select></div>
<div><label>Country</label><select id="country"></select></div>
<div><label>Platform</label><select id="platform"></select></div>
<div><label>Min Final Score</label><input id="minScore" type="number" min="0" max="100"></div>
<div><label>Symbol / Name</label><input id="search" placeholder="π.χ. ASML"></div>
<div><label>Purchase Date από</label><input id="dateFrom" type="date"></div>
<div><label>Purchase Date έως</label><input id="dateTo" type="date"></div>
<div><label>Ηλικία θέσης</label><select id="agePreset"><option value="">Όλες</option><option value="1m">1 μήνας</option><option value="3m">3 μήνες</option><option value="6m">6 μήνες</option><option value="12m">12 μήνες</option><option value="older12m">&gt; 12 μήνες</option></select></div>
</div>

<div class="kpis">
<div class="kpi"><div class="l">Portfolio Value</div><div class="v" id="kValue">—</div></div>
<div class="kpi"><div class="l">Portfolio Return</div><div class="v" id="kReturn">—</div></div>
<div class="kpi"><div class="l">VUAA Return</div><div class="v" id="kVuaa">—</div></div>
<div class="kpi"><div class="l">Excess Return</div><div class="v" id="kExcess">—</div></div>
<div class="kpi"><div class="l">Positions</div><div class="v" id="kCount">—</div></div>
<div class="kpi"><div class="l">Avg Score</div><div class="v" id="kFinal">—</div></div>
<div class="kpi"><div class="l">Quality Alerts</div><div class="v" id="kQuality">—</div></div>
</div>

<div class="grid">
<div class="card"><div class="title">Κατανομή ανά Sector</div><div id="sectorChart" class="chart"></div></div>
<div class="card"><div class="title">Κατανομή ανά Χώρα</div><div id="countryChart" class="chart"></div></div>
<div class="card"><div class="title">Μεγαλύτερες Θέσεις</div><div id="holdingsChart" class="chart"></div></div>
<div class="card"><div class="title">Συνολική Απόδοση</div><div id="returnChart" class="chart tall"></div></div>
<div class="card"><div class="title">Excess Return vs VUAA</div><div id="excessChart" class="chart tall"></div></div>
<div class="card"><div class="title">1M vs Μέση Μηνιαία</div><div id="momentumChart" class="chart tall"></div></div>

<div class="card full">
<div class="title">Υφιστάμενες Μετοχές — Final Score & Αναλυτική Αιτιολόγηση</div>
<div id="portfolioAccordion" class="accordion-list"></div>
</div>

<div class="card scoretable">
<div class="title">Υποψήφιες Μετοχές — Score & Απόσταση από 52W High</div>
<div class="tablewrap"><table><thead><tr><th>#</th><th>Μετοχή</th><th>Όνομα</th><th>Score</th><th>% κάτω από 52W High</th></tr></thead><tbody id="candidateScoreBody"></tbody></table></div>
</div>

<div class="card full">
<div class="title">AI Ανάλυση Υποψήφιων Μετοχών</div>
<div id="candidateAnalysisList" class="analysis-list"></div>
</div>

<div class="card full"><div class="title">ROIC vs FCF Yield</div><div id="fundChart" class="chart"></div></div>
<div class="card full detail"><div class="title">Αναλυτικός Πίνακας Portfolio</div><div class="tablewrap"><table><thead><tr><th>Symbol</th><th>Name</th><th>Sector</th><th>Country</th><th>Platform</th><th>Purchase Date</th><th>Market Value</th><th>Return</th><th>1M</th><th>Avg/Month</th><th>VUAA</th><th>Excess</th><th>FCF Yield</th><th>FCF Growth</th><th>ROIC</th><th>Forward Revenue</th><th>Forward EPS</th><th>Score</th></tr></thead><tbody id="tbody"></tbody></table></div></div>
</div>
</div>

<script>
const PAYLOAD=__PAYLOAD__;
const DATA=PAYLOAD.stocks||[];
const CANDIDATES=PAYLOAD.candidates||[];
const QUALITY=PAYLOAD.quality||{};
const euro=new Intl.NumberFormat("el-GR",{style:"currency",currency:"EUR"});
const pct=x=>x==null?"—":Number(x).toLocaleString("el-GR",{minimumFractionDigits:2,maximumFractionDigits:2})+"%";
const num=x=>x==null?"—":Number(x).toLocaleString("el-GR",{minimumFractionDigits:2,maximumFractionDigits:2});
const esc=x=>String(x??"").replace(/[&<>"']/g,m=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[m]));
const C={blue:"#4472C4",green:"#16845B",red:"#C43D4B",orange:"#D97706"};

function isoDate(v){if(!v)return null;const s=String(v).slice(0,10);return /^\d{4}-\d{2}-\d{2}$/.test(s)?s:null}
function uniq(f){return [...new Set(DATA.map(x=>x[f]).filter(Boolean))].sort((a,b)=>String(a).localeCompare(String(b),"el"))}
function fill(id,vals){document.getElementById(id).innerHTML='<option value="">Όλα</option>'+vals.map(v=>`<option>${v}</option>`).join("")}
fill("sector",uniq("sector"));fill("country",uniq("country"));fill("platform",uniq("platform"));

function rows(){
 const s=document.getElementById("sector").value,c=document.getElementById("country").value,p=document.getElementById("platform").value;
 const q=document.getElementById("search").value.trim().toLowerCase(),min=document.getElementById("minScore").value;
 const df=document.getElementById("dateFrom").value,dt=document.getElementById("dateTo").value,age=document.getElementById("agePreset").value;
 const now=new Date(); const cut=n=>{const d=new Date(now);d.setMonth(d.getMonth()-n);return d};
 return DATA.filter(x=>{
   if(s&&x.sector!==s||c&&x.country!==c||p&&x.platform!==p)return false;
   if(min!==""&&(x.final_score==null||Number(x.final_score)<Number(min)))return false;
   if(q&&!((x.symbol||"").toLowerCase().includes(q)||(x.name||"").toLowerCase().includes(q)))return false;
   const id=isoDate(x.purchase_date);
   if(df||dt||age){if(!id)return false;const d=new Date(id+"T00:00:00");if(df&&d<new Date(df+"T00:00:00"))return false;if(dt&&d>new Date(dt+"T23:59:59"))return false;
     if(age==="1m"&&d<cut(1)||age==="3m"&&d<cut(3)||age==="6m"&&d<cut(6)||age==="12m"&&d<cut(12)||age==="older12m"&&d>=cut(12))return false;}
   return true;
 });
}
function sum(a){return a.reduce((x,y)=>x+(Number.isFinite(Number(y))?Number(y):0),0)}
function agg(rs,f){const m=new Map();rs.forEach(x=>m.set(x[f]||"Unknown",(m.get(x[f]||"Unknown")||0)+Number(x.market_value||0)));return [...m.entries()].sort((a,b)=>b[1]-a[1])}
function lay(x={}){return Object.assign({margin:{l:60,r:25,t:15,b:55},paper_bgcolor:"#fff",plot_bgcolor:"#fff",font:{family:"Segoe UI,Arial",color:"#172033"},xaxis:{gridcolor:"#E9EDF3"},yaxis:{gridcolor:"#E9EDF3"}},x)}
function cls(v){return v==null?"":Number(v)>=0?"pos":"neg"}
function impactClass(v){const t=String(v||"").toLowerCase();if(t.includes("αρνητικό"))return "impact-negative";if(t.includes("θετικό"))return "impact-positive";return "impact-neutral"}
function statusClass(v){const t=String(v||"").toLowerCase();if(t.includes("large")||t==="error")return "status-error";if(t.includes("moderate")||t==="warning")return "status-warning";if(t.includes("consistent")||t==="ok"||t==="none")return "status-ok";return ""}
function fmtNews(x){const n=x.news_headlines||[];if(!n.length)return '<p>Δεν βρέθηκαν πρόσφατοι τίτλοι ειδήσεων στο τελευταίο refresh.</p>';return '<ul class="news-list">'+n.slice(0,3).map(h=>`<li>${esc(h.title||"")}${h.publisher?` — ${esc(h.publisher)}`:""}</li>`).join("")+'</ul>'}
function fmtMetrics(x){const arr=x.score_metric_details||[];if(!arr.length){return `<div class="metric-grid"><div class="metric-chip"><div class="m">FCF Yield</div><div class="v">${pct(x.fcf_yield)}</div></div><div class="metric-chip"><div class="m">ROIC</div><div class="v">${pct(x.roic)}</div></div><div class="metric-chip"><div class="m">FCF Growth</div><div class="v">${pct(x.fcf_growth)}</div></div><div class="metric-chip"><div class="m">Forward Revenue</div><div class="v">${pct(x.forward_revenue_growth)}</div></div><div class="metric-chip"><div class="m">Forward EPS</div><div class="v">${pct(x.forward_eps_growth)}</div></div></div>`;}return '<div class="metric-grid">'+arr.map(m=>`<div class="metric-chip"><div class="m">${esc(m.metric)}</div><div class="v ${impactClass(m.impact)}">${m.metric==="Price/Book"?num(m.value):pct(m.value)}</div><div class="m">${esc(m.impact)}</div></div>`).join("")+'</div>'}
function bullets(items,emptyText){const a=Array.isArray(items)?items:[];if(!a.length)return `<p>${esc(emptyText)}</p>`;return '<ul class="bullet-line">'+a.map(t=>`<li>${esc(t)}</li>`).join("")+'</ul>'}
function render(){
 const rs=rows(),mv=sum(rs.map(x=>x.market_value)),cost=sum(rs.map(x=>Number(x.purchase_price||0)*Number(x.quantity||0)));
 const pr=cost>0?(mv/cost-1)*100:null;
 const br=rs.filter(x=>x.vuaa_return!=null&&x.market_value!=null),bmv=sum(br.map(x=>x.market_value));
 const vr=bmv>0?sum(br.map(x=>Number(x.market_value)*Number(x.vuaa_return)))/bmv:null;
 const ex=pr!=null&&vr!=null?pr-vr:null;
 const scores=rs.map(x=>x.final_score).filter(x=>x!=null).map(Number);
 document.getElementById("kValue").textContent=euro.format(mv);
 document.getElementById("kReturn").textContent=pct(pr);
 document.getElementById("kVuaa").textContent=pct(vr);
 document.getElementById("kExcess").textContent=pct(ex);
 document.getElementById("kCount").textContent=rs.length;
 document.getElementById("kFinal").textContent=scores.length?num(sum(scores)/scores.length):"—";
 const aq=QUALITY.anomaly_detection||{};
 document.getElementById("kQuality").textContent=(QUALITY.errors||0)+"E / "+(QUALITY.warnings||0)+"W / "+(aq.total_flags||0)+"A";

 const sec=agg(rs,"sector"),cou=agg(rs,"country");
 Plotly.react("sectorChart",[{type:"pie",labels:sec.map(x=>x[0]),values:sec.map(x=>x[1]),hole:.5}],lay({margin:{l:10,r:10,t:10,b:10}}),{responsive:true,displaylogo:false});
 Plotly.react("countryChart",[{type:"pie",labels:cou.map(x=>x[0]),values:cou.map(x=>x[1]),hole:.5}],lay({margin:{l:10,r:10,t:10,b:10}}),{responsive:true,displaylogo:false});

 const h=[...rs].filter(x=>x.market_value!=null).sort((a,b)=>b.market_value-a.market_value).slice(0,15).reverse();
 Plotly.react("holdingsChart",[{type:"bar",orientation:"h",y:h.map(x=>x.symbol),x:h.map(x=>x.market_value),marker:{color:C.blue}}],lay({margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const r=[...rs].filter(x=>x.percent_change!=null).sort((a,b)=>a.percent_change-b.percent_change);
 Plotly.react("returnChart",[{type:"bar",orientation:"h",y:r.map(x=>x.symbol),x:r.map(x=>x.percent_change),marker:{color:r.map(x=>x.percent_change>=0?C.green:C.red)}}],lay({margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const e=[...rs].filter(x=>x.excess_return!=null).sort((a,b)=>a.excess_return-b.excess_return);
 Plotly.react("excessChart",[{type:"bar",orientation:"h",y:e.map(x=>x.symbol),x:e.map(x=>x.excess_return),marker:{color:e.map(x=>x.excess_return>=0?C.green:C.red)}}],lay({margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const m=[...rs].filter(x=>x.monthly_percent_change!=null||x.avg_monthly_change!=null).slice(0,25).reverse();
 Plotly.react("momentumChart",[{type:"bar",orientation:"h",name:"1M",y:m.map(x=>x.symbol),x:m.map(x=>x.monthly_percent_change),marker:{color:C.blue}},{type:"bar",orientation:"h",name:"Avg/Month",y:m.map(x=>x.symbol),x:m.map(x=>x.avg_monthly_change),marker:{color:C.orange}}],lay({barmode:"group",margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const portfolioScores=[...rs].filter(x=>x.final_score!=null).sort((a,b)=>Number(b.final_score)-Number(a.final_score));
 document.getElementById("portfolioAccordion").innerHTML=portfolioScores.map((x,i)=>{
   const source=x.secondary_source||"χωρίς 2η πηγή";
   const model=x.score_model_label||((x.financial_model)?"Financial Services":"General Fundamentals");
   const anomalyStatus=x.anomaly_count?`<span class="${statusClass(x.anomaly_severity)}">${x.anomaly_count} anomaly</span>`:'<span class="status-ok">χωρίς anomaly</span>';
   const earnings=x.earnings_date?esc(x.earnings_date):"—";
   const commentarySource=x.ai_commentary_source||"rules_fallback";
   const rationale=x.score_rationale_text||x.ai_commentary||"Δεν υπάρχει διαθέσιμη αιτιολόγηση.";
   return `<details class="stock-details"><summary class="stock-summary"><span class="stock-rank">#${i+1}</span><span class="stock-symbol">${esc(x.symbol)}</span><span class="stock-name">${esc(x.name)}</span><span class="score-badge">Score ${num(x.final_score)}</span><span class="source-badge">${esc(source)}</span></summary><div class="stock-panel"><div class="analysis-grid"><div class="subcard"><h4>Γιατί πήρε αυτό το Final Score</h4><p><b>${esc(x.score_rationale_headline||"")}</b></p><p>${esc(rationale)}</p>${fmtMetrics(x)}</div><div class="subcard"><h4>Θετικοί οδηγοί</h4>${bullets(x.score_drivers,"Δεν προέκυψε σαφής θετικός οδηγός από τα διαθέσιμα πεδία.")}<h4>Αδύνατα σημεία / κίνδυνοι</h4>${bullets(x.score_weaknesses,"Δεν προέκυψε σαφής αρνητικός οδηγός από τα διαθέσιμα πεδία.")}</div><div class="subcard"><h4>FMP / SEC ανεξάρτητος έλεγχος</h4><p>${esc(x.score_crosscheck_text||"Δεν υπάρχει διαθέσιμο cross-check.")}</p><p><b>Κατάσταση:</b> <span class="${statusClass(x.fundamental_crosscheck_status)}">${esc(x.fundamental_crosscheck_status||"—")}</span>${x.fundamental_discrepancy_pct!=null?` • Απόκλιση ${pct(x.fundamental_discrepancy_pct)}`:""}</p><p><b>Μοντέλο score:</b> ${esc(model)}</p></div><div class="subcard"><h4>News / Earnings / Data Quality</h4>${fmtNews(x)}<p><b>Earnings:</b> ${earnings} • <b>Sentiment:</b> ${esc(x.news_sentiment||"—")}</p><p><b>Data check:</b> ${anomalyStatus}${x.anomaly_summary?` — ${esc(x.anomaly_summary)}`:""}</p></div><div class="subcard" style="grid-column:1/-1"><h4>AI Commentary</h4><p>${esc(x.ai_commentary||"Δεν υπάρχει διαθέσιμο commentary.")}</p><p class="analysis-meta">Πηγή commentary: ${esc(commentarySource)}${x.ai_commentary_error?` • OpenAI error: ${esc(x.ai_commentary_error)}`:""}</p></div></div></div></details>`;
 }).join("");

 const candidateScores=[...CANDIDATES].filter(x=>x.final_score!=null).sort((a,b)=>(Number(a.rank||999)-Number(b.rank||999))||(Number(b.final_score)-Number(a.final_score)));
 document.getElementById("candidateScoreBody").innerHTML=candidateScores.map((x,i)=>`<tr><td>${x.rank||i+1}</td><td><b>${esc(x.symbol)}</b></td><td>${esc(x.name)}</td><td><b>${num(x.final_score)}</b></td><td>${pct(x.discount_from_52w_high_pct)}</td></tr>`).join("");
 document.getElementById("candidateAnalysisList").innerHTML=candidateScores.map((x,i)=>{
   const earnings=x.earnings_date?` • Earnings ${esc(x.earnings_date)}`:"";
   const source=x.ai_commentary_source?` • ${esc(x.ai_commentary_source)}`:"";
   return `<article class="analysis-item"><h3>${x.rank||i+1}. ${esc(x.symbol)} — ${esc(x.name)}</h3><div class="analysis-meta">Score ${num(x.final_score)} • ${pct(x.discount_from_52w_high_pct)} κάτω από 52W High${earnings}${source}</div><p>${esc(x.dashboard_analysis||x.ai_commentary||"Δεν υπάρχει διαθέσιμη σύντομη ανάλυση για αυτή την υποψήφια.")}</p></article>`;
 }).join("");

 const f=rs.filter(x=>x.roic!=null&&x.fcf_yield!=null);
 Plotly.react("fundChart",[{type:"scatter",mode:"markers+text",textposition:"top center",text:f.map(x=>x.symbol),x:f.map(x=>x.roic),y:f.map(x=>x.fcf_yield),marker:{size:10,color:C.blue}}],lay({xaxis:{title:"ROIC (%)",gridcolor:"#E9EDF3"},yaxis:{title:"FCF Yield (%)",gridcolor:"#E9EDF3"}}),{responsive:true,displaylogo:false});

 document.getElementById("tbody").innerHTML=[...rs].sort((a,b)=>(b.market_value||0)-(a.market_value||0)).map(x=>`<tr><td><b>${x.symbol||""}</b></td><td>${x.name||""}</td><td>${x.sector||""}</td><td>${x.country||""}</td><td>${x.platform||""}</td><td>${String(x.purchase_date||"").slice(0,10)}</td><td>${x.market_value==null?"—":euro.format(x.market_value)}</td><td class="${cls(x.percent_change)}">${pct(x.percent_change)}</td><td class="${cls(x.monthly_percent_change)}">${pct(x.monthly_percent_change)}</td><td class="${cls(x.avg_monthly_change)}">${pct(x.avg_monthly_change)}</td><td>${pct(x.vuaa_return)}</td><td class="${cls(x.excess_return)}">${pct(x.excess_return)}</td><td>${pct(x.fcf_yield)}</td><td>${pct(x.fcf_growth)}</td><td>${pct(x.roic)}</td><td>${pct(x.forward_revenue_growth)}</td><td>${pct(x.forward_eps_growth)}</td><td><b>${num(x.final_score)}</b></td></tr>`).join("");
}

["sector","country","platform","minScore","dateFrom","dateTo","agePreset"].forEach(id=>document.getElementById(id).addEventListener("change",render));
document.getElementById("search").addEventListener("input",render);
render();
</script>
</body>
</html>'''


def generate_portfolio_dashboard(
    stocks,
    quality_report=None,
    candidates=None,
    json_path=DASHBOARD_JSON_FILE,
    html_path=DASHBOARD_HTML_FILE,
):
    json_path = Path(json_path)
    html_path = Path(html_path)

    payload = {
        "generated_at": datetime.now().isoformat(
            timespec="seconds"
        ),
        "stocks": _serializable_stocks(
            stocks
        ),
        "candidates": (
            _serializable_stocks(candidates)
            if candidates is not None
            else _load_candidates_from_excel()
        ),
        "quality": quality_report or {},
    }

    json_path.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
            default=_json_default,
        ),
        encoding="utf-8",
    )

    compact_payload = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        default=_json_default,
    )

    html = HTML_TEMPLATE.replace(
        "__PAYLOAD__",
        compact_payload,
    )

    html_path.write_text(
        html,
        encoding="utf-8",
    )

    return {
        "json_file": json_path,
        "html_file": html_path,
    }


def refresh_dashboard_candidates(candidates=None):
    if not DASHBOARD_JSON_FILE.exists():
        return None

    try:
        previous = json.loads(
            DASHBOARD_JSON_FILE.read_text(
                encoding="utf-8"
            )
        )
    except Exception:
        return None

    return generate_portfolio_dashboard(
        stocks=previous.get(
            "stocks",
            [],
        ),
        quality_report=previous.get(
            "quality",
            {},
        ),
        candidates=candidates,
    )
