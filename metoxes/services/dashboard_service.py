from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]
DASHBOARD_JSON_FILE = BASE_DIR / "portfolio_dashboard_data.json"
DASHBOARD_HTML_FILE = BASE_DIR / "portfolio_dashboard.html"
PORTFOLIO_EXCEL_FILE = BASE_DIR / "portfolio.xlsx"
CANDIDATES_SHEET_NAME = "Candidates"

HIDDEN_SCORE_FIELDS = {
    "relative_score",
    "absolute_score",
    "absolute_coverage",
    "score_mode",
}


def _json_default(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()

    raise TypeError(
        f"Object of type {type(value).__name__} "
        "is not JSON serializable"
    )


def _public_stock_view(stock):
    return {
        key: value
        for key, value in stock.items()
        if key not in HIDDEN_SCORE_FIELDS
        and not key.startswith("absolute_")
        and not key.startswith("relative_")
        and "_absolute_score" not in key
        and "_relative_score" not in key
    }


def _serializable_stocks(stocks):
    public_rows = [
        _public_stock_view(stock)
        for stock in stocks or []
    ]

    return json.loads(
        json.dumps(
            public_rows,
            ensure_ascii=False,
            default=_json_default,
        )
    )


def _load_candidates_from_excel():
    if not PORTFOLIO_EXCEL_FILE.exists():
        return []

    try:
        workbook = load_workbook(
            PORTFOLIO_EXCEL_FILE,
            data_only=True,
            read_only=True,
        )
    except Exception:
        return []

    try:
        if CANDIDATES_SHEET_NAME not in workbook.sheetnames:
            return []

        sheet = workbook[CANDIDATES_SHEET_NAME]
        headers = {}

        for cell in sheet[1]:
            if cell.value is None:
                continue

            headers[
                str(cell.value).strip().lower()
            ] = cell.column

        wanted = (
            "rank",
            "symbol",
            "name",
            "sector",
            "country",
            "currency",
            "current_price",
            "high_52w",
            "price_25pct_below_high",
            "drop_from_52w_high_amount",
            "discount_from_52w_high_pct",
            "meets_25pct_discount",
            "already_in_portfolio",
            "fcf_yield",
            "fcf_growth",
            "roic",
            "forward_revenue_growth",
            "forward_eps_growth",
            "final_score",
            "market_cap",
            "validation_status",
            "validation_source",
            "researched_at",
            "source_url",
            "dashboard_analysis",
        )

        rows = []

        for row_number in range(2, sheet.max_row + 1):
            symbol_col = headers.get("symbol")

            if symbol_col is None:
                break

            symbol = sheet.cell(
                row=row_number,
                column=symbol_col,
            ).value

            if not symbol:
                continue

            row = {}

            for field in wanted:
                column = headers.get(field)

                if column is not None:
                    row[field] = sheet.cell(
                        row=row_number,
                        column=column,
                    ).value

            rows.append(row)

        # Even when only the compact Candidates sheet is available, build a
        # deterministic score explanation from the stored screening metrics.
        # Rich FMP/SEC/news/AI fields are preserved separately from the previous
        # dashboard JSON when available.
        try:
            from metoxes.services.score_explanation_service import (
                build_score_explanation,
            )
            for row in rows:
                if str(row.get("sector") or "") == "Financial Services":
                    row.setdefault(
                        "score_rationale_headline",
                        "Η υποψήφια αξιολογείται με ειδικό Financial Services μοντέλο."
                    )
                    row.setdefault(
                        "score_rationale_text",
                        row.get("dashboard_analysis") or ""
                    )
                    row.setdefault(
                        "score_model_label",
                        "Financial Services"
                    )
                    row.setdefault("score_drivers", [])
                    row.setdefault("score_weaknesses", [])
                    row.setdefault("score_metric_details", [])
                else:
                    for key, value in build_score_explanation(row).items():
                        row.setdefault(key, value)
        except Exception:
            pass

        return rows
    finally:
        workbook.close()


def _load_candidates_preserving_enrichment(json_path=DASHBOARD_JSON_FILE):
    """
    Keep rich candidate research fields across normal portfolio refreshes.

    /stocks/research-candidates writes enriched candidate rows (score rationale,
    FMP/SEC, news, earnings, anomalies, AI commentary). A later
    /stocks/update-excel used to rebuild candidates only from the compact Excel
    sheet and silently discard those fields. Merge the current Excel snapshot
    with the previous dashboard candidate payload so current visible values win
    while research enrichment survives.
    """
    excel_rows = _load_candidates_from_excel()
    json_path = Path(json_path)

    previous_rows = []
    if json_path.exists():
        try:
            payload = json.loads(json_path.read_text(encoding="utf-8"))
            previous_rows = payload.get("candidates") or []
        except Exception:
            previous_rows = []

    previous_by_symbol = {
        str(row.get("symbol") or "").strip().upper(): dict(row)
        for row in previous_rows
        if isinstance(row, dict) and row.get("symbol")
    }

    if not excel_rows:
        return list(previous_by_symbol.values())

    merged = []
    for row in excel_rows:
        symbol = str(row.get("symbol") or "").strip().upper()
        combined = dict(previous_by_symbol.get(symbol) or {})
        combined.update(row)
        merged.append(combined)

    return merged


HTML_TEMPLATE = r'''<!DOCTYPE html>
<html lang="el">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>METOXES2 Portfolio Dashboard</title>
<script src="https://cdn.plot.ly/plotly-3.3.1.min.js"></script>
<style>
:root{--blue:#4472C4;--blue2:#2F5597;--bg:#F4F7FB;--card:#fff;--text:#172033;--muted:#697386;--border:#DCE3ED;--green:#16845B;--red:#C43D4B;--orange:#D97706}
*{box-sizing:border-box}body{margin:0;font-family:Segoe UI,Arial,sans-serif;background:var(--bg);color:var(--text)}
.header{background:linear-gradient(135deg,var(--blue2),var(--blue));color:#fff;padding:22px 28px}.header h1{margin:0;font-size:26px}.header p{margin:7px 0 0;opacity:.88}
.wrap{max-width:1550px;margin:auto;padding:18px}.filters{display:grid;grid-template-columns:repeat(8,minmax(135px,1fr));gap:10px;background:var(--card);padding:14px;border:1px solid var(--border);border-radius:12px;margin-bottom:14px}
label{font-size:12px;color:var(--muted);font-weight:600;display:block;margin-bottom:5px}select,input{width:100%;padding:8px;border:1px solid var(--border);border-radius:8px;background:#fff}
.kpis{display:grid;grid-template-columns:repeat(7,1fr);gap:10px;margin-bottom:14px}.kpi,.card{background:var(--card);border:1px solid var(--border);border-radius:12px}.kpi{padding:14px}.kpi .l{font-size:11px;color:var(--muted);text-transform:uppercase}.kpi .v{font-size:22px;font-weight:700;margin-top:7px}
.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}.card{padding:8px}.full{grid-column:1/-1}.title{font-weight:700;padding:10px 12px 0}.chart{height:410px}.tall{height:510px}
.tablewrap{max-height:560px;overflow:auto;margin:8px;border:1px solid var(--border);border-radius:8px}
table{border-collapse:collapse;width:100%;font-size:12px}th{position:sticky;top:0;background:var(--blue);color:#fff;padding:8px;white-space:nowrap}td{padding:8px;border-bottom:1px solid #E9EDF3;white-space:nowrap}.pos{color:var(--green);font-weight:600}.neg{color:var(--red);font-weight:600}
.scoretable table{min-width:0}.scoretable .tablewrap{max-height:520px}.detail table{min-width:1750px}
.analysis-list{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px;padding:10px}.analysis-item{border:1px solid var(--border);border-radius:10px;padding:14px;background:#FAFCFF}.analysis-item h3{font-size:14px;margin:0 0 7px}.analysis-meta{font-size:12px;color:var(--muted);margin-bottom:8px}.analysis-item p{font-size:13px;line-height:1.55;margin:0}
.master-details{border:1px solid var(--border);border-radius:12px;background:#fff;overflow:hidden}.master-summary{display:flex;align-items:center;justify-content:space-between;gap:12px;padding:14px 16px;cursor:pointer;list-style:none;font-weight:800;font-size:16px;background:#F8FAFD}.master-summary::-webkit-details-marker{display:none}.master-summary::after{content:"▼";color:var(--blue2);font-size:17px;font-weight:900;transition:transform .18s ease;flex:0 0 auto}.master-details[open]>.master-summary::after{transform:rotate(180deg)}.master-summary:hover{background:#F1F5FB}.master-content{border-top:1px solid var(--border);padding:0}.accordion-list{display:grid;gap:9px;padding:10px}.stock-details{border:1px solid var(--border);border-radius:11px;background:#fff;overflow:hidden}.stock-details[open]{box-shadow:0 5px 18px rgba(31,78,121,.08)}.stock-summary{display:grid;grid-template-columns:48px 110px minmax(180px,1fr) 120px 125px 30px;gap:10px;align-items:center;padding:12px 14px;cursor:pointer;list-style:none}.stock-summary::-webkit-details-marker{display:none}.stock-summary::after{content:"▼";justify-self:end;color:var(--blue2);font-size:15px;font-weight:900;transition:transform .18s ease}.stock-details[open]>.stock-summary::after{transform:rotate(180deg)}.stock-summary:hover{background:#F8FAFD}.stock-rank{color:var(--muted);font-size:12px}.stock-symbol{font-weight:800}.stock-name{font-weight:600}.score-badge{display:inline-flex;justify-content:center;align-items:center;border-radius:999px;padding:6px 10px;font-weight:800;background:#EAF1FB;color:var(--blue2)}.source-badge{font-size:11px;text-align:center;border:1px solid var(--border);border-radius:999px;padding:5px 8px;color:var(--muted)}.stock-panel{border-top:1px solid var(--border);padding:14px;background:#FBFCFE}.analysis-grid{display:grid;grid-template-columns:1.2fr 1fr;gap:12px}.subcard{border:1px solid #E5EAF1;border-radius:9px;background:#fff;padding:12px}.subcard h4{margin:0 0 8px;font-size:13px}.subcard p{font-size:12.5px;line-height:1.55;margin:4px 0}.metric-grid{display:grid;grid-template-columns:repeat(5,minmax(110px,1fr));gap:8px;margin-top:10px}.metric-chip{border:1px solid #E5EAF1;border-radius:8px;padding:8px;background:#fff}.metric-chip .m{font-size:11px;color:var(--muted)}.metric-chip .v{font-size:14px;font-weight:700;margin-top:3px}.impact-positive{color:var(--green)}.impact-negative{color:var(--red)}.impact-neutral{color:var(--orange)}.bullet-line{margin:4px 0;padding-left:16px}.bullet-line li{margin:3px 0;font-size:12.5px;line-height:1.45}.news-list{margin:5px 0 0;padding-left:16px}.news-list li{margin:3px 0;font-size:12.5px;line-height:1.4}.candidate-details .stock-summary{grid-template-columns:48px 110px minmax(180px,1fr) 120px 150px 30px}.context-grid{display:grid;grid-template-columns:repeat(4,minmax(130px,1fr));gap:8px;margin-top:8px}.context-chip{border:1px solid #E5EAF1;border-radius:8px;padding:9px;background:#fff}.context-chip .m{font-size:11px;color:var(--muted)}.context-chip .v{font-size:13px;font-weight:700;margin-top:3px}.status-warning{color:var(--orange);font-weight:700}.status-error{color:var(--red);font-weight:700}.status-ok{color:var(--green);font-weight:700}.model-summary{padding:12px 14px;display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px}.model-box{border:1px solid var(--border);border-radius:9px;padding:10px;background:#FAFCFF}.model-box .m{font-size:11px;color:var(--muted)}.model-box .v{font-size:14px;font-weight:700;margin-top:4px}.peer-list{font-size:12px;color:var(--muted);margin-top:6px}
@media(max-width:1200px){.filters{grid-template-columns:repeat(4,1fr)}.kpis{grid-template-columns:repeat(4,1fr)}.model-summary{grid-template-columns:repeat(2,1fr)}}@media(max-width:900px){.filters,.kpis,.grid,.analysis-list,.analysis-grid{grid-template-columns:1fr}.stock-summary,.candidate-details .stock-summary{grid-template-columns:34px 80px 1fr 88px 24px}.stock-summary .source-badge{display:none}.metric-grid,.context-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.full{grid-column:auto}}

.quarter-card{background:#f2f7f5!important;border:1px solid #d8e7e0!important}.quarter-card .stock-panel,.quarter-card .master-summary{background:#f2f7f5}.quarter-stats{display:flex;flex-wrap:wrap;gap:12px;margin:12px 0}.quarter-stats .subcard{flex:1;min-width:190px}.quarter-review-table{width:100%;font-size:12px}.quarter-note{font-size:12px;color:#526475;line-height:1.6}.candidate-details .stock-summary{grid-template-columns:35px 85px minmax(140px,1fr) 95px 90px 75px 65px 100px 120px 20px}.candidate-label{display:block;font-size:10px;font-weight:400;color:#5b6c7e}.country-badge{text-align:center;font-size:12px;font-weight:700;color:#355b50;background:#edf6f2;border:1px solid #d3e8df;border-radius:999px;padding:6px 8px}.beta-badge{text-align:center;font-weight:700}.investing-check{text-align:center}.investing-check input{accent-color:#4d806a;pointer-events:none}.low52{text-align:center;font-size:12px}
@media(max-width:1100px){.candidate-details .stock-summary{grid-template-columns:30px 70px minmax(120px,1fr) 80px 75px 65px 55px 80px 90px 20px;gap:5px;padding:10px}.candidateAccordionWrap{overflow-x:auto}#candidateAccordion{min-width:880px}.candidate-details .source-badge{display:block!important}}
.saved-investing .stock-summary{grid-template-columns:36px minmax(220px,1fr) 100px 65px 105px 125px 20px}.saved-investing .stock-name{white-space:normal}.saved-investing .stock-details{min-width:800px}.saved-investing .stock-panel{white-space:normal}
#portfolioMaster>.master-summary{background:#edf4fb;border-left:5px solid #6c9bc7;color:#244b70}#portfolioMaster>.master-summary:hover{background:#e4eff9}
#candidateMaster>.master-summary{background:#f1eef9;border-left:5px solid #8d79b8;color:#514174}#candidateMaster>.master-summary:hover{background:#e9e4f5}
#modelMaster>.master-summary{background:#fff5e6;border-left:5px solid #d5a04f;color:#75501c}#modelMaster>.master-summary:hover{background:#ffefd7}
.quarter-card .master-summary{border-left:5px solid #6eaa91;color:#355b50}
</style>
</head>
<body>
<div class="header"><h1>METOXES2 Portfolio Dashboard</h1><p>Portfolio • υποψήφιες μετοχές • συνολικό score • 52-week high • data quality</p></div>
<div class="wrap">

<div class="filters">
<div><label>Sector</label><select id="sector"></select></div>
<div><label>Country</label><select id="country"></select></div>
<div><label>Platform</label><select id="platform"></select></div>
<div><label>Min Final Score</label><input id="minScore" type="number" min="0" max="100"></div>
<div><label>Symbol / Name</label><input id="search" placeholder="π.χ. ASML"></div>
<div><label>Purchase Date από</label><input id="dateFrom" type="date"></div>
<div><label>Purchase Date έως</label><input id="dateTo" type="date"></div>
<div><label>Ηλικία θέσης</label><select id="agePreset"><option value="">Όλες</option><option value="1m">1 μήνας</option><option value="3m">3 μήνες</option><option value="6m">6 μήνες</option><option value="12m">12 μήνες</option><option value="older12m">&gt; 12 μήνες</option></select></div>
</div>

<div class="kpis">
<div class="kpi"><div class="l">Portfolio Value</div><div class="v" id="kValue">—</div></div>
<div class="kpi"><div class="l">Portfolio Return</div><div class="v" id="kReturn">—</div></div>
<div class="kpi"><div class="l">VUAA Return</div><div class="v" id="kVuaa">—</div></div>
<div class="kpi"><div class="l">Excess Return</div><div class="v" id="kExcess">—</div></div>
<div class="kpi"><div class="l">Positions</div><div class="v" id="kCount">—</div></div>
<div class="kpi"><div class="l">Avg Score</div><div class="v" id="kFinal">—</div></div>
<div class="kpi"><div class="l">Quality Alerts</div><div class="v" id="kQuality">—</div></div>
</div>

<div class="grid">
<div class="card"><div class="title">Κατανομή ανά Sector</div><div id="sectorChart" class="chart"></div></div>
<div class="card"><div class="title">Κατανομή ανά Χώρα</div><div id="countryChart" class="chart"></div></div>
<div class="card"><div class="title">Μεγαλύτερες Θέσεις</div><div id="holdingsChart" class="chart"></div></div>
<div class="card"><div class="title">Συνολική Απόδοση</div><div id="returnChart" class="chart tall"></div></div>
<div class="card"><div class="title">Excess Return vs VUAA</div><div id="excessChart" class="chart tall"></div></div>
<div class="card"><div class="title">1M vs Μέση Μηνιαία</div><div id="momentumChart" class="chart tall"></div></div>

<div class="card full">
<details class="master-details" id="portfolioMaster">
<summary class="master-summary">Υφιστάμενες Μετοχές — Final Score & Αναλυτική Αιτιολόγηση</summary>
<div class="master-content">
<div id="portfolioAccordion" class="accordion-list"></div>
</div>
</details>
</div>

<div class="card full quarter-card"><details open><summary class="master-summary">Αγορές τελευταίου τριμήνου — πρόοδος και σύγκριση</summary><div id="quarterReview" class="stock-panel"></div></details></div>
<div class="card full quarter-card"><details open><summary class="master-summary">Αγορές τελευταίου εξαμήνου — μέση επίδοση τελευταίων 3 μηνών</summary><div id="sixMonthReview" class="stock-panel"></div></details></div>
<div class="card full">
<details class="master-details" id="candidateMaster">
<summary class="master-summary">Υποψήφιες Μετοχές — Final Score & Αναλυτική Αξιολόγηση</summary>
<div class="master-content">
<div id="candidateAnalysisList" style="display:none"></div>
<div class="candidateAccordionWrap"><div id="candidateAccordion" class="accordion-list"></div></div>
</div>
</details>
</div>

<div class="card full">
<details class="master-details" id="modelMaster">
<summary class="master-summary">Model Intelligence — Backtest & Learning</summary>
<div class="master-content"><div id="modelIntelligenceSummary" class="model-summary"></div></div>
</details>
</div>

<div class="card full"><div class="title">ROIC vs FCF Yield</div><div id="fundChart" class="chart"></div></div>
<div class="card full detail"><div class="title">Αναλυτικός Πίνακας Portfolio</div><div class="tablewrap"><table><thead><tr><th>Symbol</th><th>Name</th><th>Sector</th><th>Country</th><th>Platform</th><th>Purchase Date</th><th>Market Value</th><th>Return</th><th>1M</th><th>Avg/Month</th><th>VUAA</th><th>Excess</th><th>FCF Yield</th><th>FCF Growth</th><th>ROIC</th><th>Forward Revenue</th><th>Forward EPS</th><th>Score</th></tr></thead><tbody id="tbody"></tbody></table></div></div>
</div>
</div>

<script>
const PAYLOAD=__PAYLOAD__;
const DATA=PAYLOAD.stocks||[];
const CANDIDATES=PAYLOAD.candidates||[];
const QUALITY=PAYLOAD.quality||{};
const MODEL=PAYLOAD.model_intelligence||{};
const euro=new Intl.NumberFormat("el-GR",{style:"currency",currency:"EUR"});
const pct=x=>x==null?"—":Number(x).toLocaleString("el-GR",{minimumFractionDigits:2,maximumFractionDigits:2})+"%";
const num=x=>x==null?"—":Number(x).toLocaleString("el-GR",{minimumFractionDigits:2,maximumFractionDigits:2});
const esc=x=>String(x??"").replace(/[&<>"']/g,m=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[m]));
const C={blue:"#4472C4",green:"#16845B",red:"#C43D4B",orange:"#D97706"};

function safeLink(url,label){try{const u=new URL(url);return ['https:','http:'].includes(u.protocol)?`<a href="${esc(u.href)}" target="_blank" rel="noopener noreferrer">${esc(label)}</a>`:''}catch{return ''}}

function renderSixMonthReview(){
 const q=PAYLOAD.six_month_review;if(!q)return;
 const stat=(title,g)=>`<div class="subcard"><h4>${title}</h4><p><b>Μέση μηνιαία επίδοση τελευταίων 3 μηνών</b></p><p>${g.positions} θέσεις / ${g.securities} tickers · με πλήρες ιστορικό: ${g.valid_monthly}</p><p>Μέσος: <b class="${cls(g.mean_monthly_pct)}">${pct(g.mean_monthly_pct)}</b> · Διάμεσος: ${pct(g.median_monthly_pct)}</p><p>Με θετικό μέσο 3Μ: ${pct(g.positive_pct)}</p></div>`;
 const table=rows=>`<div class="tablewrap"><table class="quarter-review-table"><thead><tr><th>Μετοχή / broker</th><th>Αγορά</th><th>Μηνιαίες μεταβολές (παλαιότερη → νεότερη)</th><th>Μέση μηνιαία 3Μ</th><th>Συνολική 3Μ</th><th>Μέση τιμή κλεισίματος 3Μ</th><th>Τεκμηρίωση</th></tr></thead><tbody>${rows.map(x=>`<tr><td><b>${esc(x.symbol)}</b><br>${esc(x.platform)}</td><td>${esc(x.purchase_date)}</td><td>${(x.history.monthly_returns||[]).map(v=>pct(v)).join(' / ')||'—'}</td><td class="${cls(x.monthly_percent_change)}">${pct(x.monthly_percent_change)}${x.includes_pre_purchase?' *':''}</td><td>${pct(x.history.total_3m_pct)}</td><td>${num(x.history.mean_close_3m)} ${esc(x.currency||'')}</td><td><details><summary>Περίοδος / πηγή</summary><p>${esc(x.history.source||'Δεν υπάρχει πλήρες, ενημερωμένο ιστορικό τιμών.')}</p><p>Κλεισίματα: ${esc((x.history.price_dates||[]).join(' → '))}</p></details></td></tr>`).join('')}</tbody></table></div>`;
 document.getElementById('sixMonthReview').innerHTML=`<p>Νέες αγορές: ${esc(q.cutoff)} έως ${esc(q.as_of)} · παλαιότερες: πριν από ${esc(q.cutoff)}.</p><p><b>Κοινή περίοδος μέτρησης: ${esc(q.performance_start)} έως ${esc(q.as_of)}.</b></p><div class="quarter-stats">${stat('Νέες αγορές εξαμήνου',q.recent)}${stat('Παλαιότερες αγορές',q.older)}<div class="subcard"><h4>Διαφορά μέσης μηνιαίας επίδοσης 3Μ</h4><b class="${cls(q.difference_pp)}">${num(q.difference_pp)} ποσοστιαίες μονάδες</b><p>Νέες μείον παλαιότερες</p></div></div><p class="quarter-note">${esc(q.method)}</p><details open><summary>Νέες αγορές — αναλυτικά</summary>${table(q.rows)}</details><details><summary>Παλαιότερες αγορές — αναλυτικά</summary>${table(q.older_rows)}</details><p class="quarter-note">* Περιλαμβάνονται ημέρες πριν από την αγορά. Παύλα = ανεπαρκές ιστορικό, δεν μετρά ως μηδέν. Η ενότητα αφορά όλο το portfolio.</p>`;
}
function renderSavedInvesting(){
 const groups=PAYLOAD.investing_saved_lists||[];
 document.getElementById('savedInvestingLists').innerHTML=groups.map(g=>`<div class="card saved-investing"><details><summary class="master-summary">Investing — ${g.country==='Greece'?'Ελλάδα':'ΗΠΑ'} · Υποτιμημένες μετοχές (${g.rows.length})</summary><div class="stock-panel"><p><b>Αποθηκευμένη λίστα από το αρχείο σου.</b> Παραλαβή: ${esc(g.snapshot_date)}. Οι τιμές/αξιολογήσεις είναι όπως στο αρχείο· δεν ανανεώνονται ζωντανά από το Investing.</p><p>${safeLink(g.source_url,'Αρχική λίστα Investing')} · ${esc(g.filename)}</p><div class="tablewrap">${g.rows.map(x=>{const f=x.fields||{},q=x.quote||{};return `<details class="stock-details"><summary class="stock-summary"><span class="stock-rank">#${x.rank}</span><span class="stock-name">${esc(x.name)}</span><span class="score-badge"><span class="candidate-label">Final Score</span>${num(x.final_score)}</span><span class="beta-badge"><span class="candidate-label">Beta</span>${num(q.beta)}</span><span class="low52"><span class="candidate-label">Χαμηλό 52W</span>${num(q.low_52w)} ${esc(q.currency||'')}</span><span class="source-badge"><span class="candidate-label">Κάτω από υψηλό 52W</span>${pct(x.discount_from_52w_high_pct)}</span></summary><div class="stock-panel"><div class="analysis-grid"><div class="subcard"><h4>Final Score / υψηλό 52W</h4><p><b>Final Score: ${num(x.final_score)}</b> / 100</p><p>${esc(x.score_details?.scope||'Εκκρεμεί αντιστοίχιση και επαρκή θεμελιώδη δεδομένα για τον υπολογισμό με το μοντέλο του project.')}</p><p>FCF yield: ${pct(x.score_details?.fcf_yield)} · FCF growth: ${pct(x.score_details?.fcf_growth)} · ROIC: ${pct(x.score_details?.roic)}</p><p>Forward revenue growth: ${pct(x.score_details?.forward_revenue_growth)} · Forward EPS growth: ${pct(x.score_details?.forward_eps_growth)}</p><p>${esc(x.score_details?.score_rationale_text||'')}</p><p>Score: ${esc(x.score_details?.checked_at||'—')}</p><p>Τρέχουσα τιμή: ${num(x.current_price)} · Υψηλό 52W: ${num(x.high_52w)} · Κάτω από υψηλό: ${pct(x.discount_from_52w_high_pct)}</p><p>Ανάκτηση τιμών: ${esc(x.market_checked_at||'—')}</p><p class="quarter-note">Υπολογισμός: (1 − τρέχουσα τιμή / υψηλό 52W) × 100. Χρησιμοποιούνται τιμή και υψηλό από το ίδιο σύνολο δεδομένων. Το upside Investing παραμένει διαφορετικό μέγεθος.</p></div><div class="subcard"><h4>Αξιολογήσεις Investing</h4>${Object.entries(f).filter(([k])=>k!=='Όνομα').map(([k,v])=>`<p><b>${esc(k)}:</b> ${esc(v)}</p>`).join('')}<p>${safeLink(x.source_url,'Μετοχή στο Investing')}</p></div><div class="subcard"><h4>Beta / χαμηλό 52 εβδομάδων</h4>${q.status==='matched'?`<p>Ticker: ${esc(q.symbol)} · ${esc(q.currency)}</p><p>Beta (5Y Monthly): ${num(q.beta)} · Χαμηλό 52W: ${num(q.low_52w)}</p><p>${safeLink(q.source_url,'Yahoo Finance')} · Ανάκτηση: ${esc(q.checked_at)}</p>`:'<p>Δεν υπάρχει ακόμη επιβεβαιωμένη αντιστοίχιση και ανάκτηση. Τα κενά παραμένουν παύλες.</p>'}<p class="quarter-note">Οι ημερομηνίες Yahoo και αποθηκευμένου Investing μπορεί να διαφέρουν. Το beta είναι ιστορική εκτίμηση. Το upside του Investing δεν είναι Final Score ή εγγύηση απόδοσης.</p></div></div></div></details>`}).join('')}</div></div></details></div>`).join('');
}

function renderUSValue(){
 const q=PAYLOAD.us_value_screen;if(!q)return;
 const rows=q.rows||[];
 document.getElementById('usValueScreen').innerHTML=`<p><b>Ενημερωμένα στοιχεία: ${q.fresh_count}/${q.universe_size} σύμβολα · πέρασαν τα φίλτρα: ${rows.length}.</b> ${q.status==='complete'?'Ο έλεγχος ολοκληρώθηκε.':'Μερική κάλυψη — η λίστα μπορεί να αλλάξει όταν ολοκληρωθεί η ανάκτηση.'}</p><p>Τελευταία επιτυχής ανάκτηση: ${esc(q.checked_at||'Δεν έχει ολοκληρωθεί')} · Αυτόματη ενημέρωση μαζί με το dashboard, cache 24 ωρών.</p><details><summary>Κριτήρια και τρόπος βαθμολόγησης</summary><p class="quarter-note">${esc(q.method)}</p><p class="quarter-note">Τα θεμελιώδη μεγέθη αφορούν τις περιόδους/εκτιμήσεις του παρόχου. Η ημερομηνία ανάκτησης δεν είναι ημερομηνία οικονομικών καταστάσεων. Δεν έχει γίνει ακόμη backtesting αυτού του φίλτρου.</p><p class="quarter-note">Σύμβολα δείγματος: ${esc((q.universe||[]).join(', '))}</p></details>${!rows.length?`<p>${q.fresh_count?'Καμία εταιρεία με επαρκή στοιχεία δεν πέρασε όλα τα φίλτρα.':'Δεν ολοκληρώθηκε η άντληση στοιχείων. Η επόμενη ενημέρωση θα ξαναδοκιμάσει αυτόματα.'} Δεν εμφανίζονται υποθετικές βαθμολογίες.</p>`:''}${Object.keys(q.last_errors||{}).length?`<p class="quarter-note">Αποτυχίες ανάκτησης: ${Object.keys(q.last_errors).length}. Η επανάληψη για αυτά τα σύμβολα επιτρέπεται μετά από 15 λεπτά.</p>`:''}${rows.map(x=>`<details class="stock-details candidate-details"><summary class="stock-summary"><span class="stock-rank">#${x.rank}</span><span class="stock-symbol">${esc(x.symbol)}</span><span class="stock-name">${esc(x.name)}${x.already_in_portfolio?' · ήδη στο portfolio':''}</span><span class="score-badge"><span class="candidate-label">Score φίλτρου</span>${num(x.value_score)}</span><span class="beta-badge"><span class="candidate-label">Beta</span>${num(x.beta)}</span><span class="source-badge"><span class="candidate-label">FCF yield</span>${pct(x.fcf_yield_pct)}</span><span class="low52"><span class="candidate-label">Χαμηλό 52W</span>${num(x.low_52w)} USD</span><span class="source-badge"><span class="candidate-label">Από χαμηλό 52W</span>${pct(x.distance_from_52w_low_pct)}</span></summary><div class="stock-panel"><div class="analysis-grid"><div class="subcard"><h4>Γιατί πέρασε το φίλτρο</h4><p>${esc(x.rationale)}</p><p>Κλάδος: ${esc(x.sector)}. Δείγμα σύγκρισης: ${esc(x.peer_symbols.join(', '))}.</p><p>Χαμηλότερο forward P/E από διάμεσο κλάδου: ${pct(x.pe_discount_pct)}. Αυτό δεν είναι αναμενόμενη άνοδος τιμής.</p></div><div class="subcard"><h4>Ανάλυση score / 100</h4><p>Σχετικό P/E: ${num(x.components.relative_pe)}/30</p><p>FCF yield: ${num(x.components.fcf_yield)}/30</p><p>ROE: ${num(x.components.roe)}/20</p><p>Καθαρό περιθώριο: ${num(x.components.profit_margin)}/10</p><p>Ανάπτυξη εσόδων: ${num(x.components.revenue_growth)}/10</p></div><div class="subcard"><h4>Τιμή / Beta / 52W</h4><p>Τιμή: ${num(x.current_price)} USD · Υψηλό 52W: ${num(x.high_52w)} USD</p><p>Beta (5Y Monthly): ${num(x.beta)}. Ιστορική ευαισθησία στον δείκτη του παρόχου, όχι πρόβλεψη. ${x.beta==null?'Μη διαθέσιμο.':Math.abs(x.beta)>1.5?'Αυξημένη ευαισθησία στην αγορά.':''}</p><p>${safeLink(x.source_url,'Yahoo Finance — στοιχεία εταιρείας')} · Ανάκτηση: ${esc(x.checked_at)}</p></div><div class="subcard"><h4>Κίνδυνοι / περιορισμοί</h4><ul>${x.risks.map(v=>`<li>${esc(v)}</li>`).join('')}</ul></div></div></div></details>`).join('')}`;
}
function renderQuarter(){
 const q=PAYLOAD.quarter_review;if(!q)return;
 const stat=(title,g)=>`<div class="subcard"><h4>${title}</h4><p><b>Επίδοση τελευταίου μήνα (L / 1Μ)</b></p><p>${g.positions} θέσεις / ${g.securities} tickers · έγκυρη L: ${g.valid_monthly}</p><p>Μέση μεταβολή τελευταίου μήνα: <b class="${cls(g.mean_monthly_pct)}">${pct(g.mean_monthly_pct)}</b> · Διάμεσος τελευταίου μήνα: ${pct(g.median_monthly_pct)}</p><p>Θετικές τελευταίου μήνα: ${pct(g.positive_pct)}</p><p>Συνολική απόδοση από αγορά: ${pct(g.since_purchase_pct)}</p></div>`;
 document.getElementById('quarterReview').innerHTML=`<p>Αγορές ${esc(q.cutoff)} έως ${esc(q.as_of)}. Σύγκριση με αγορές πριν από ${esc(q.cutoff)}.</p><div class="quarter-stats">${stat('Νέες αγορές',q.recent)}${stat('Παλαιότερες αγορές',q.older)}<div class="subcard"><h4>Διαφορά μέσης μηνιαίας επίδοσης</h4><b class="${cls(q.difference_pp)}">${num(q.difference_pp)} ποσοστιαίες μονάδες</b><p>Νέες μείον παλαιότερες</p></div></div><p class="quarter-note">${esc(q.method)}</p><p class="quarter-note">Έρευνα γενικών ροών ειδήσεων από Yahoo Finance, CNBC, MarketWatch, Investing.com · ${esc(q.research?.status||'')} · ${num(q.research?.elapsed_seconds)} sec · όριο 170 sec. Τελευταίος έλεγχος: ${esc(q.research?.checked_at||'δεν εκτελέστηκε')}. Το όριο αφορά μόνο την πρόσθετη έρευνα V21. Οι ροές καλύπτουν τους διαθέσιμους πρόσφατους τίτλους, όχι πλήρες αρχείο μηνός.</p><p class="quarter-note">Η ενότητα αφορά όλο το portfolio, ανεξάρτητα από τα επάνω φίλτρα.</p><div class="tablewrap"><table class="quarter-review-table"><thead><tr><th>Μετοχή / broker</th><th>Αγορά</th><th>Τιμή αγοράς → τώρα</th><th>Από αγορά</th><th>L / 1M</th><th>Ειδήσεις / πιθανή εξήγηση</th></tr></thead><tbody>${q.rows.map(x=>`<tr><td><b>${esc(x.symbol)}</b><br>${esc(x.platform)}</td><td>${esc(x.purchase_date)}</td><td>${num(x.purchase_price)} → ${num(x.current_price)}</td><td class="${cls(x.percent_change)}">${pct(x.percent_change)}</td><td class="${cls(x.monthly_percent_change)}">${pct(x.monthly_percent_change)}${x.monthly_includes_pre_purchase?' *':''}</td><td><details><summary>Ανάλυση και πηγές (${(x.news?.articles||[]).length})</summary><p>${esc(x.news?.summary||'Δεν έχει ολοκληρωθεί έρευνα για αυτή τη μετοχή.')}</p><ul>${(x.news?.articles||[]).map(a=>`<li>${safeLink(a.url,a.title)} — ${esc(a.source)} · ${esc(a.published_at)}</li>`).join('')}</ul><small>${esc(JSON.stringify(x.news?.source_status||{}))}</small></details></td></tr>`).join('')}</tbody></table></div><p class="quarter-note">* Κατοχή κάτω από 30 ημέρες: η L μπορεί να περιέχει μεταβολή πριν από την αγορά. Τιμές όπως καταγράφονται στο portfolio.</p>`;
}

function isoDate(v){if(!v)return null;const s=String(v).slice(0,10);return /^\d{4}-\d{2}-\d{2}$/.test(s)?s:null}
function uniq(f){return [...new Set(DATA.map(x=>x[f]).filter(Boolean))].sort((a,b)=>String(a).localeCompare(String(b),"el"))}
function fill(id,vals){document.getElementById(id).innerHTML='<option value="">Όλα</option>'+vals.map(v=>`<option>${v}</option>`).join("")}
fill("sector",uniq("sector"));fill("country",uniq("country"));fill("platform",uniq("platform"));

function rows(){
 const s=document.getElementById("sector").value,c=document.getElementById("country").value,p=document.getElementById("platform").value;
 const q=document.getElementById("search").value.trim().toLowerCase(),min=document.getElementById("minScore").value;
 const df=document.getElementById("dateFrom").value,dt=document.getElementById("dateTo").value,age=document.getElementById("agePreset").value;
 const now=new Date(); const cut=n=>{const d=new Date(now);d.setMonth(d.getMonth()-n);return d};
 return DATA.filter(x=>{
   if(s&&x.sector!==s||c&&x.country!==c||p&&x.platform!==p)return false;
   if(min!==""&&(x.final_score==null||Number(x.final_score)<Number(min)))return false;
   if(q&&!((x.symbol||"").toLowerCase().includes(q)||(x.name||"").toLowerCase().includes(q)))return false;
   const id=isoDate(x.purchase_date);
   if(df||dt||age){if(!id)return false;const d=new Date(id+"T00:00:00");if(df&&d<new Date(df+"T00:00:00"))return false;if(dt&&d>new Date(dt+"T23:59:59"))return false;
     if(age==="1m"&&d<cut(1)||age==="3m"&&d<cut(3)||age==="6m"&&d<cut(6)||age==="12m"&&d<cut(12)||age==="older12m"&&d>=cut(12))return false;}
   return true;
 });
}
function sum(a){return a.reduce((x,y)=>x+(Number.isFinite(Number(y))?Number(y):0),0)}
function agg(rs,f){const m=new Map();rs.forEach(x=>m.set(x[f]||"Unknown",(m.get(x[f]||"Unknown")||0)+Number(x.market_value||0)));return [...m.entries()].sort((a,b)=>b[1]-a[1])}
function lay(x={}){return Object.assign({margin:{l:60,r:25,t:15,b:55},paper_bgcolor:"#fff",plot_bgcolor:"#fff",font:{family:"Segoe UI,Arial",color:"#172033"},xaxis:{gridcolor:"#E9EDF3"},yaxis:{gridcolor:"#E9EDF3"}},x)}
function cls(v){return v==null?"":Number(v)>=0?"pos":"neg"}
function impactClass(v){const t=String(v||"").toLowerCase();if(t.includes("αρνητικό"))return "impact-negative";if(t.includes("θετικό"))return "impact-positive";return "impact-neutral"}
function statusClass(v){const t=String(v||"").toLowerCase();if(t.includes("large")||t==="error")return "status-error";if(t.includes("moderate")||t==="warning")return "status-warning";if(t.includes("consistent")||t==="ok"||t==="none")return "status-ok";return ""}
function fmtNews(x){const n=x.news_headlines||[];if(!n.length)return '<p>Δεν βρέθηκαν πρόσφατοι τίτλοι ειδήσεων στο τελευταίο refresh.</p>';return '<ul class="news-list">'+n.slice(0,3).map(h=>`<li>${esc(h.title||"")}${h.publisher?` — ${esc(h.publisher)}`:""}</li>`).join("")+'</ul>'}
function fmtMetrics(x){const arr=x.score_metric_details||[];if(!arr.length){return `<div class="metric-grid"><div class="metric-chip"><div class="m">FCF Yield</div><div class="v">${pct(x.fcf_yield)}</div></div><div class="metric-chip"><div class="m">ROIC</div><div class="v">${pct(x.roic)}</div></div><div class="metric-chip"><div class="m">FCF Growth</div><div class="v">${pct(x.fcf_growth)}</div></div><div class="metric-chip"><div class="m">Forward Revenue</div><div class="v">${pct(x.forward_revenue_growth)}</div></div><div class="metric-chip"><div class="m">Forward EPS</div><div class="v">${pct(x.forward_eps_growth)}</div></div></div>`;}return '<div class="metric-grid">'+arr.map(m=>`<div class="metric-chip"><div class="m">${esc(m.metric)}</div><div class="v ${impactClass(m.impact)}">${m.metric==="Price/Book"?num(m.value):pct(m.value)}</div><div class="m">${esc(m.impact)}</div></div>`).join("")+'</div>'}
function bullets(items,emptyText){const a=Array.isArray(items)?items:[];if(!a.length)return `<p>${esc(emptyText)}</p>`;return '<ul class="bullet-line">'+a.map(t=>`<li>${esc(t)}</li>`).join("")+'</ul>'}
function fmtPeer(x){const syms=Array.isArray(x.peer_symbols)?x.peer_symbols:[];return `<p><b>${esc(x.peer_summary||"Δεν υπάρχουν αρκετά peer δεδομένα.")}</b></p>${bullets(x.peer_strengths,"Δεν προέκυψε σαφές peer πλεονέκτημα.")}${bullets(x.peer_weaknesses,"Δεν προέκυψε σαφής peer αδυναμία.")}<p class="peer-list">Peers: ${syms.length?syms.map(esc).join(", "):"—"} • Πηγή: ${esc(x.peer_source||"fallback/none")}</p>`}
function fmtEventImpact(x){const score=x.event_impact_score==null?"—":`${Number(x.event_impact_score)>=0?"+":""}${num(x.event_impact_score)}`;return `<p><b>Impact:</b> <span class="${Number(x.event_impact_score||0)>15?"impact-positive":Number(x.event_impact_score||0)<-15?"impact-negative":"impact-neutral"}">${score}/100 • ${esc(x.event_impact_label||"neutral")}</span></p><p>${esc(x.event_impact_summary||"Δεν υπάρχουν αρκετά event δεδομένα.")}</p><p class="analysis-meta">Confidence: ${esc(x.event_impact_confidence||"low")} • Source: ${esc(x.event_impact_source||"rules_fallback")}</p>`}
function fmtScenario(x){return `<div class="context-grid"><div class="context-chip"><div class="m">EPS -15pp</div><div class="v">${num(x.scenario_eps_minus_15_score)}</div><div class="m">Δ ${x.scenario_eps_minus_15_delta==null?"—":num(x.scenario_eps_minus_15_delta)}</div></div><div class="context-chip"><div class="m">Growth slowdown</div><div class="v">${num(x.scenario_growth_slowdown_score)}</div><div class="m">Δ ${x.scenario_growth_slowdown_delta==null?"—":num(x.scenario_growth_slowdown_delta)}</div></div><div class="context-chip"><div class="m">Bear case</div><div class="v">${num(x.scenario_bear_score)}</div><div class="m">Δ ${x.scenario_bear_delta==null?"—":num(x.scenario_bear_delta)}</div></div><div class="context-chip"><div class="m">Sensitivity</div><div class="v">${esc(x.scenario_risk_label||"—")}</div></div></div><p>${esc(x.scenario_summary||"")}</p>`}
function fmtModelIntelligence(){const q=MODEL.quartile_test||{},ic=MODEL.final_score_predictive_ic||{};return `<div class="model-box"><div class="m">Backtest status</div><div class="v">${esc(MODEL.status||"not_run_yet")}</div><div class="m">${Number(MODEL.evaluated_samples||0)} matured samples • ${Number(MODEL.horizon_days||0)||"—"}d horizon</div></div><div class="model-box"><div class="m">Final Score predictive IC</div><div class="v">${ic.directional_ic==null?"—":num(ic.directional_ic)}</div><div class="m">Spearman rank correlation</div></div><div class="model-box"><div class="m">Top-minus-bottom quartile</div><div class="v">${q.top_minus_bottom_spread_pct==null?"—":pct(q.top_minus_bottom_spread_pct)}</div><div class="m">Top hit rate ${q.top_quartile_positive_hit_rate_pct==null?"—":pct(q.top_quartile_positive_hit_rate_pct)}</div></div><div class="model-box"><div class="m">Learning policy</div><div class="v">Recommendations only</div><div class="m">No silent live-weight changes</div></div>`}
function render(){
 const rs=rows(),mv=sum(rs.map(x=>x.market_value)),cost=sum(rs.map(x=>Number(x.purchase_price||0)*Number(x.quantity||0)));
 const pr=cost>0?(mv/cost-1)*100:null;
 const br=rs.filter(x=>x.vuaa_return!=null&&x.market_value!=null),bmv=sum(br.map(x=>x.market_value));
 const vr=bmv>0?sum(br.map(x=>Number(x.market_value)*Number(x.vuaa_return)))/bmv:null;
 const ex=pr!=null&&vr!=null?pr-vr:null;
 const scores=rs.map(x=>x.final_score).filter(x=>x!=null).map(Number);
 document.getElementById("kValue").textContent=euro.format(mv);
 document.getElementById("kReturn").textContent=pct(pr);
 document.getElementById("kVuaa").textContent=pct(vr);
 document.getElementById("kExcess").textContent=pct(ex);
 document.getElementById("kCount").textContent=rs.length;
 document.getElementById("kFinal").textContent=scores.length?num(sum(scores)/scores.length):"—";
 const aq=QUALITY.anomaly_detection||{};
 document.getElementById("kQuality").textContent=(QUALITY.errors||0)+"E / "+(QUALITY.warnings||0)+"W / "+(aq.total_flags||0)+"A";

 const sec=agg(rs,"sector"),cou=agg(rs,"country");
 Plotly.react("sectorChart",[{type:"pie",labels:sec.map(x=>x[0]),values:sec.map(x=>x[1]),hole:.5}],lay({margin:{l:10,r:10,t:10,b:10}}),{responsive:true,displaylogo:false});
 Plotly.react("countryChart",[{type:"pie",labels:cou.map(x=>x[0]),values:cou.map(x=>x[1]),hole:.5}],lay({margin:{l:10,r:10,t:10,b:10}}),{responsive:true,displaylogo:false});

 const h=[...rs].filter(x=>x.market_value!=null).sort((a,b)=>b.market_value-a.market_value).slice(0,15).reverse();
 Plotly.react("holdingsChart",[{type:"bar",orientation:"h",y:h.map(x=>x.symbol),x:h.map(x=>x.market_value),marker:{color:C.blue}}],lay({margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const r=[...rs].filter(x=>x.percent_change!=null).sort((a,b)=>a.percent_change-b.percent_change);
 Plotly.react("returnChart",[{type:"bar",orientation:"h",y:r.map(x=>x.symbol),x:r.map(x=>x.percent_change),marker:{color:r.map(x=>x.percent_change>=0?C.green:C.red)}}],lay({margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const e=[...rs].filter(x=>x.excess_return!=null).sort((a,b)=>a.excess_return-b.excess_return);
 Plotly.react("excessChart",[{type:"bar",orientation:"h",y:e.map(x=>x.symbol),x:e.map(x=>x.excess_return),marker:{color:e.map(x=>x.excess_return>=0?C.green:C.red)}}],lay({margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const m=[...rs].filter(x=>x.monthly_percent_change!=null||x.avg_monthly_change!=null).slice(0,25).reverse();
 Plotly.react("momentumChart",[{type:"bar",orientation:"h",name:"1M",y:m.map(x=>x.symbol),x:m.map(x=>x.monthly_percent_change),marker:{color:C.blue}},{type:"bar",orientation:"h",name:"Avg/Month",y:m.map(x=>x.symbol),x:m.map(x=>x.avg_monthly_change),marker:{color:C.orange}}],lay({barmode:"group",margin:{l:80,r:25,t:15,b:55}}),{responsive:true,displaylogo:false});

 const portfolioScores=[...rs].filter(x=>x.final_score!=null).sort((a,b)=>Number(b.final_score)-Number(a.final_score));
 document.getElementById("portfolioAccordion").innerHTML=portfolioScores.map((x,i)=>{
   const source=x.secondary_source||"χωρίς 2η πηγή";
   const model=x.score_model_label||((x.financial_model)?"Financial Services":"General Fundamentals");
   const anomalyStatus=x.anomaly_count?`<span class="${statusClass(x.anomaly_severity)}">${x.anomaly_count} anomaly</span>`:'<span class="status-ok">χωρίς anomaly</span>';
   const earnings=x.earnings_date?esc(x.earnings_date):"—";
   const commentarySource=x.ai_commentary_source||"rules_fallback";
   const rationale=x.score_rationale_text||x.ai_commentary||"Δεν υπάρχει διαθέσιμη αιτιολόγηση.";
   return `<details class="stock-details"><summary class="stock-summary"><span class="stock-rank">#${i+1}</span><span class="stock-symbol">${esc(x.symbol)}</span><span class="stock-name">${esc(x.name)}</span><span class="score-badge">Score ${num(x.final_score)}</span><span class="source-badge">${esc(source)}</span></summary><div class="stock-panel"><div class="analysis-grid"><div class="subcard"><h4>Γιατί πήρε αυτό το Final Score</h4><p><b>${esc(x.score_rationale_headline||"")}</b></p><p>${esc(rationale)}</p>${fmtMetrics(x)}</div><div class="subcard"><h4>Θετικοί οδηγοί</h4>${bullets(x.score_drivers,"Δεν προέκυψε σαφής θετικός οδηγός από τα διαθέσιμα πεδία.")}<h4>Αδύνατα σημεία / κίνδυνοι</h4>${bullets(x.score_weaknesses,"Δεν προέκυψε σαφής αρνητικός οδηγός από τα διαθέσιμα πεδία.")}</div><div class="subcard"><h4>FMP / SEC ανεξάρτητος έλεγχος</h4><p>${esc(x.score_crosscheck_text||"Δεν υπάρχει διαθέσιμο cross-check.")}</p><p><b>Κατάσταση:</b> <span class="${statusClass(x.fundamental_crosscheck_status)}">${esc(x.fundamental_crosscheck_status||"—")}</span>${x.fundamental_discrepancy_pct!=null?` • Απόκλιση ${pct(x.fundamental_discrepancy_pct)}`:""}</p><p><b>Μοντέλο score:</b> ${esc(model)}</p></div><div class="subcard"><h4>News / Earnings / Data Quality</h4>${fmtNews(x)}<p><b>Earnings:</b> ${earnings} • <b>Sentiment:</b> ${esc(x.news_sentiment||"—")}</p><p><b>Data check:</b> ${anomalyStatus}${x.anomaly_summary?` — ${esc(x.anomaly_summary)}`:""}</p></div><div class="subcard"><h4>Πραγματικοί ανταγωνιστές / Peers</h4>${fmtPeer(x)}</div><div class="subcard"><h4>News & Earnings Impact</h4>${fmtEventImpact(x)}</div><div class="subcard" style="grid-column:1/-1"><h4>Scenario Analysis</h4>${fmtScenario(x)}</div><div class="subcard" style="grid-column:1/-1"><h4>AI Commentary</h4><p>${esc(x.ai_commentary||"Δεν υπάρχει διαθέσιμο commentary.")}</p><p class="analysis-meta">Πηγή commentary: ${esc(commentarySource)}${x.ai_commentary_error?` • OpenAI error: ${esc(x.ai_commentary_error)}`:""}</p></div></div></div></details>`;
 }).join("");

 const candidateScores=[...CANDIDATES].filter(x=>x.final_score!=null).sort((a,b)=>(Number(a.rank||999)-Number(b.rank||999))||(Number(b.final_score)-Number(a.final_score)));
 document.getElementById("candidateAccordion").innerHTML=candidateScores.map((x,i)=>{
   const source=x.secondary_source||"χωρίς 2η πηγή";
   const model=x.score_model_label||((x.financial_model)?"Financial Services":"General Fundamentals");
   const earnings=x.earnings_date?esc(x.earnings_date):"—";
   const anomalyStatus=x.anomaly_count?`<span class="${statusClass(x.anomaly_severity)}">${x.anomaly_count} anomaly</span>`:'<span class="status-ok">χωρίς anomaly</span>';
   const commentarySource=x.ai_commentary_source||"rules_fallback";
   const narrative=x.dashboard_analysis||x.ai_commentary||x.score_rationale_text||"Δεν υπάρχει διαθέσιμη αξιολόγηση.";
   const validation=x.validation_status||x.secondary_validation_status||"—";
   return `<details class="stock-details candidate-details"><summary class="stock-summary"><span class="stock-rank">#${x.rank||i+1}</span><span class="stock-symbol">${esc(x.symbol)}</span><span class="stock-name">${esc(x.name)}</span><span class="score-badge">Score ${num(x.final_score)}</span><span class="beta-badge" title="${esc(x.beta_source||'Μη διαθέσιμο')}"><span class="candidate-label">Beta</span>${num(x.beta)}${x.beta_status==='stale'?' *':''}</span><span class="investing-check" title="${esc((x.investing_status||'not_checked')+' • '+(x.investing_checked_at||''))}"><span class="candidate-label">Investing</span>${x.investing_undervalued==null?'—':`<input type="checkbox" disabled ${x.investing_undervalued?'checked':''} aria-label="${x.investing_undervalued?'Στη λίστα Investing':'Δεν εμφανίζεται στη λίστα Investing'}">`}</span><span class="low52"><span class="candidate-label">52W χαμηλό</span>${num(x.low_52w)}<small> ${esc(x.currency||'')}</small></span><span class="source-badge">${pct(x.discount_from_52w_high_pct)} κάτω από υψηλό</span></summary><div class="stock-panel"><div class="analysis-grid"><div class="subcard"><h4>Beta / Investing</h4><p>${esc(x.beta_assessment||'Μη διαθέσιμο beta.')}</p><p class="analysis-meta">${esc(x.beta_source||'')} • ${esc(x.beta_checked_at||'')} • ${esc(x.beta_status||'')} ${safeLink(x.beta_source_url,'Πηγή beta')}</p><p>Investing: ${esc(x.investing_status||'not_checked')} — ${esc(x.investing_scope||'')} Έλεγχος: ${esc(x.investing_checked_at||'—')}. ${safeLink(x.investing_source_url,'Άνοιγμα λίστας')}</p><p>Απόσταση από χαμηλό 52 εβδομάδων: ${pct(x.distance_from_52w_low_pct)}.</p><h4>Συνολική αξιολόγηση υποψήφιας</h4><p><b>${esc(x.score_rationale_headline||"")}</b></p><p>${esc(narrative)}</p>${fmtMetrics(x)}</div><div class="subcard"><h4>Θετικοί οδηγοί</h4>${bullets(x.score_drivers,"Δεν προέκυψε σαφής θετικός οδηγός από τα διαθέσιμα πεδία.")}<h4>Αδύνατα σημεία / κίνδυνοι</h4>${bullets(x.score_weaknesses,"Δεν προέκυψε σαφής αρνητικός οδηγός από τα διαθέσιμα πεδία.")}</div><div class="subcard"><h4>Τιμή / 52W High / Validation</h4><div class="context-grid"><div class="context-chip"><div class="m">Τρέχουσα τιμή</div><div class="v">${x.current_price==null?"—":num(x.current_price)}</div></div><div class="context-chip"><div class="m">52W High</div><div class="v">${x.high_52w==null?"—":num(x.high_52w)}</div></div><div class="context-chip"><div class="m">% κάτω από 52W</div><div class="v">${pct(x.discount_from_52w_high_pct)}</div></div><div class="context-chip"><div class="m">Validation</div><div class="v">${esc(validation)}</div></div></div><p><b>Μοντέλο score:</b> ${esc(model)}</p></div><div class="subcard"><h4>FMP / SEC ανεξάρτητος έλεγχος</h4><p>${esc(x.score_crosscheck_text||"Δεν υπάρχει διαθέσιμο cross-check.")}</p><p><b>Πηγή:</b> ${esc(source)} • <b>Κατάσταση:</b> <span class="${statusClass(x.fundamental_crosscheck_status)}">${esc(x.fundamental_crosscheck_status||"—")}</span>${x.fundamental_discrepancy_pct!=null?` • Απόκλιση ${pct(x.fundamental_discrepancy_pct)}`:""}</p></div><div class="subcard"><h4>News / Earnings / Data Quality</h4>${fmtNews(x)}<p><b>Earnings:</b> ${earnings} • <b>Sentiment:</b> ${esc(x.news_sentiment||"—")}</p><p><b>Data check:</b> ${anomalyStatus}${x.anomaly_summary?` — ${esc(x.anomaly_summary)}`:""}</p></div><div class="subcard"><h4>Πραγματικοί ανταγωνιστές / Peers</h4>${fmtPeer(x)}</div><div class="subcard"><h4>News & Earnings Impact</h4>${fmtEventImpact(x)}</div><div class="subcard" style="grid-column:1/-1"><h4>Scenario Analysis</h4>${fmtScenario(x)}</div><div class="subcard" style="grid-column:1/-1"><h4>AI Commentary / τελική επεξήγηση</h4><p>${esc(x.ai_commentary||x.score_rationale_text||"Δεν υπάρχει διαθέσιμο commentary.")}</p><p class="analysis-meta">Πηγή commentary: ${esc(commentarySource)}${x.ai_commentary_error?` • OpenAI error: ${esc(x.ai_commentary_error)}`:""}</p></div></div></div></details>`;
 }).join("");
 document.querySelectorAll("#candidateAccordion .candidate-details").forEach((row,i)=>{
   const x=candidateScores[i]||{};
   const country=x.country||x.market_country||x.investing_country||"—";
   const badge=document.createElement("span");
   badge.className="country-badge";
   badge.innerHTML=`<span class="candidate-label">Χώρα</span>${esc(country)}`;
   const score=row.querySelector(".score-badge");
   if(score)score.before(badge);
 });

 document.getElementById("modelIntelligenceSummary").innerHTML=fmtModelIntelligence();

 const f=rs.filter(x=>x.roic!=null&&x.fcf_yield!=null);
 Plotly.react("fundChart",[{type:"scatter",mode:"markers+text",textposition:"top center",text:f.map(x=>x.symbol),x:f.map(x=>x.roic),y:f.map(x=>x.fcf_yield),marker:{size:10,color:C.blue}}],lay({xaxis:{title:"ROIC (%)",gridcolor:"#E9EDF3"},yaxis:{title:"FCF Yield (%)",gridcolor:"#E9EDF3"}}),{responsive:true,displaylogo:false});

 document.getElementById("tbody").innerHTML=[...rs].sort((a,b)=>(b.market_value||0)-(a.market_value||0)).map(x=>`<tr><td><b>${x.symbol||""}</b></td><td>${x.name||""}</td><td>${x.sector||""}</td><td>${x.country||""}</td><td>${x.platform||""}</td><td>${String(x.purchase_date||"").slice(0,10)}</td><td>${x.market_value==null?"—":euro.format(x.market_value)}</td><td class="${cls(x.percent_change)}">${pct(x.percent_change)}</td><td class="${cls(x.monthly_percent_change)}">${pct(x.monthly_percent_change)}</td><td class="${cls(x.avg_monthly_change)}">${pct(x.avg_monthly_change)}</td><td>${pct(x.vuaa_return)}</td><td class="${cls(x.excess_return)}">${pct(x.excess_return)}</td><td>${pct(x.fcf_yield)}</td><td>${pct(x.fcf_growth)}</td><td>${pct(x.roic)}</td><td>${pct(x.forward_revenue_growth)}</td><td>${pct(x.forward_eps_growth)}</td><td><b>${num(x.final_score)}</b></td></tr>`).join("");
}

["sector","country","platform","minScore","dateFrom","dateTo","agePreset"].forEach(id=>document.getElementById(id).addEventListener("change",render));
document.getElementById("search").addEventListener("input",render);
renderQuarter();
renderSixMonthReview();
render();
</script>
</body>
</html>'''


def generate_portfolio_dashboard(
    stocks,
    quality_report=None,
    candidates=None,
    json_path=DASHBOARD_JSON_FILE,
    html_path=DASHBOARD_HTML_FILE,
):
    json_path = Path(json_path)
    html_path = Path(html_path)

    payload = {
        "generated_at": datetime.now().isoformat(
            timespec="seconds"
        ),
        "stocks": _serializable_stocks(
            stocks
        ),
        "candidates": (
            _serializable_stocks(candidates)
            if candidates is not None
            else _serializable_stocks(
                _load_candidates_preserving_enrichment(
                    json_path=json_path
                )
            )
        ),
        "quality": quality_report or {},
        "model_intelligence": (
            __import__(
                "metoxes.services.model_backtest_service",
                fromlist=["load_latest_backtest_report"],
            ).load_latest_backtest_report()
        ),
    }

    from .portfolio_review_service import enrich_payload
    payload = enrich_payload(payload, root=BASE_DIR)
    if candidates is not None:
        enriched = {r.get("symbol"): r for r in payload["candidates"]}
        for candidate in candidates:
            candidate.update(enriched.get(candidate.get("symbol"), {}))

    json_path.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
            default=_json_default,
        ),
        encoding="utf-8",
    )

    compact_payload = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        default=_json_default,
    )

    html = HTML_TEMPLATE.replace(
        "__PAYLOAD__",
        compact_payload.replace("<", "\\u003c"),
    )

    html_path.write_text(
        html,
        encoding="utf-8",
    )

    return {
        "json_file": json_path,
        "html_file": html_path,
    }


def refresh_dashboard_candidates(candidates=None):
    if not DASHBOARD_JSON_FILE.exists():
        return None

    try:
        previous = json.loads(
            DASHBOARD_JSON_FILE.read_text(
                encoding="utf-8"
            )
        )
    except Exception:
        return None

    return generate_portfolio_dashboard(
        stocks=previous.get(
            "stocks",
            [],
        ),
        quality_report=previous.get(
            "quality",
            {},
        ),
        candidates=candidates,
    )
