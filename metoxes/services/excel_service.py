from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ==========================================================
# ΑΡΧΕΙΑ
# ==========================================================

BASE_DIR = Path(__file__).resolve().parents[2]

EXCEL_FILE = BASE_DIR / "portfolio.xlsx"
OUTPUT_FILE = BASE_DIR / "portfolio_updated.xlsx"


# ==========================================================
# ΣΤΗΛΕΣ EXCEL
# ==========================================================

HEADERS = [
    "symbol",
    "name",
    "sector",
    "country",
    "platform",
    "purchase_date",
    "purchase_price",
    "quantity",
    "current_price",
    "market_value",
    "percent_change",
    "monthly_percent_change",
    "portfolio_weight",
    "vuaa_return",
    "excess_return",
    "fcf_yield",
    "fcf_growth",
    "roic",
    "final_score",
]


# ==========================================================
# ΑΦΑΙΡΕΣΗ ΠΑΛΙΩΝ EXCEL TABLES
# ==========================================================

def remove_existing_tables(sheet):

    table_names = list(
        sheet.tables.keys()
    )

    for table_name in table_names:

        del sheet.tables[
            table_name
        ]


# ==========================================================
# ΚΑΘΑΡΙΣΜΟΣ ΕΠΙΠΛΕΟΝ ΣΤΗΛΩΝ
# ==========================================================

def remove_extra_columns(sheet):

    wanted_columns = len(
        HEADERS
    )

    if sheet.max_column > wanted_columns:

        columns_to_delete = (
            sheet.max_column
            - wanted_columns
        )

        sheet.delete_cols(
            wanted_columns + 1,
            columns_to_delete
        )


# ==========================================================
# ΣΩΣΤΑ HEADERS
# ==========================================================

def write_correct_headers(sheet):

    for column, header in enumerate(
        HEADERS,
        start=1
    ):

        sheet.cell(
            row=1,
            column=column
        ).value = header


# ==========================================================
# EXCEL TABLE
# ==========================================================

def create_stock_table(sheet):

    # Χρειαζόμαστε τουλάχιστον
    # header + 1 μετοχή.
    if sheet.max_row < 2:
        return

    last_column = get_column_letter(
        len(HEADERS)
    )

    table_reference = (
        f"A1:{last_column}{sheet.max_row}"
    )

    # ------------------------------------------------------
    # Δημιουργία Table
    # ------------------------------------------------------

    table = Table(
        displayName="StockSummary",
        ref=table_reference
    )

    # ------------------------------------------------------
    # ΜΠΛΕ EXCEL TABLE STYLE
    #
    # Σκούρο μπλε header
    # Λευκά γράμματα
    # Λευκή / ανοιχτό μπλε εναλλαγή
    # ------------------------------------------------------

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    # ΠΡΟΣΟΧΗ:
    # table = instance
    # Table = class
    table.tableStyleInfo = style

    sheet.add_table(
        table
    )


# ==========================================================
# ΜΟΡΦΟΠΟΙΗΣΗ EXCEL
# ==========================================================

def format_portfolio_sheet(sheet):

    # ------------------------------------------------------
    # Freeze header
    # ------------------------------------------------------

    sheet.freeze_panes = "A2"

    # ------------------------------------------------------
    # HEADER
    #
    # Μπλε #4472C4
    # Λευκά γράμματα
    # ------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="4472C4"
    )

    for column_number in range(
        1,
        len(HEADERS) + 1
    ):

        cell = sheet.cell(
            row=1,
            column=column_number
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    sheet.row_dimensions[1].height = 24

    # ------------------------------------------------------
    # HEADER LOOKUP
    # ------------------------------------------------------

    headers = {}

    for cell in sheet[1]:

        if cell.value:

            header = str(
                cell.value
            ).strip().lower()

            headers[
                header
            ] = cell.column

    # ------------------------------------------------------
    # ΧΡΗΜΑΤΙΚΕΣ ΣΤΗΛΕΣ
    # ------------------------------------------------------

    money_fields = [
        "purchase_price",
        "current_price",
        "market_value",
    ]

    for field in money_fields:

        if field not in headers:
            continue

        column = headers[
            field
        ]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            sheet.cell(
                row=row,
                column=column
            ).number_format = '#,##0.00 "€"'

    # ------------------------------------------------------
    # ΠΟΣΟΣΤΑ
    #
    # Οι τιμές είναι ήδη:
    #
    # 2.03
    #
    # και όχι:
    #
    # 0.0203
    # ------------------------------------------------------

    percent_fields = [
        "percent_change",
        "monthly_percent_change",
        "portfolio_weight",
        "vuaa_return",
        "excess_return",
        "fcf_yield",
        "fcf_growth",
        "roic",
    ]

    for field in percent_fields:

        if field not in headers:
            continue

        column = headers[
            field
        ]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            sheet.cell(
                row=row,
                column=column
            ).number_format = '0.00"%"'

    # ------------------------------------------------------
    # FINAL SCORE
    # ------------------------------------------------------

    if "final_score" in headers:

        column = headers[
            "final_score"
        ]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            sheet.cell(
                row=row,
                column=column
            ).number_format = "0.00"

    # ------------------------------------------------------
    # QUANTITY
    # ------------------------------------------------------

    if "quantity" in headers:

        column = headers[
            "quantity"
        ]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            sheet.cell(
                row=row,
                column=column
            ).number_format = "0.########"

    # ------------------------------------------------------
    # PURCHASE DATE
    # ------------------------------------------------------

    if "purchase_date" in headers:

        column = headers[
            "purchase_date"
        ]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            sheet.cell(
                row=row,
                column=column
            ).alignment = Alignment(
                horizontal="center"
            )

    # ------------------------------------------------------
    # ΑΥΤΟΜΑΤΟ ΠΛΑΤΟΣ ΣΤΗΛΩΝ
    # ------------------------------------------------------

    for column_number in range(
        1,
        len(HEADERS) + 1
    ):

        max_length = 0

        for row in range(
            1,
            sheet.max_row + 1
        ):

            cell = sheet.cell(
                row=row,
                column=column_number
            )

            if cell.value is None:
                continue

            length = len(
                str(cell.value)
            )

            if length > max_length:
                max_length = length

        column_letter = get_column_letter(
            column_number
        )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            25
        )

    # ------------------------------------------------------
    # EXCEL TABLE
    # ------------------------------------------------------

    create_stock_table(
        sheet
    )


# ==========================================================
# ΔΙΑΒΑΣΜΑ portfolio.xlsx
# ==========================================================

def get_excel_stocks():

    if not EXCEL_FILE.exists():

        raise FileNotFoundError(
            f"Δεν βρέθηκε το Excel: {EXCEL_FILE}"
        )

    workbook = load_workbook(
        EXCEL_FILE,
        data_only=True
    )

    sheet = workbook.worksheets[0]

    # ------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------

    headers = {}

    for cell in sheet[1]:

        if cell.value:

            header = str(
                cell.value
            ).strip().lower()

            headers[
                header
            ] = cell.column

    if "symbol" not in headers:

        workbook.close()

        raise ValueError(
            "Δεν βρέθηκε στήλη 'symbol' στο Excel"
        )

    # ------------------------------------------------------
    # STOCKS
    # ------------------------------------------------------

    stocks = []

    for row in range(
        2,
        sheet.max_row + 1
    ):

        symbol = sheet.cell(
            row=row,
            column=headers["symbol"]
        ).value

        if not symbol:
            continue

        stock = {
            "symbol": str(
                symbol
            ).strip()
        }

        if "platform" in headers:

            stock["platform"] = sheet.cell(
                row=row,
                column=headers["platform"]
            ).value

        if "sector" in headers:

            stock["sector"] = sheet.cell(
                row=row,
                column=headers["sector"]
            ).value

        if "country" in headers:

            stock["country"] = sheet.cell(
                row=row,
                column=headers["country"]
            ).value

        stocks.append(
            stock
        )

    workbook.close()

    return stocks


# ==========================================================
# SYMBOL LOOKUP
# ==========================================================

def get_symbol_lookup():

    stocks = get_excel_stocks()

    lookup = {}

    for stock in stocks:

        symbol = stock[
            "symbol"
        ].strip()

        # Πλήρες ticker
        lookup[
            symbol.upper()
        ] = symbol

        # Βασικό ticker
        base_symbol = symbol.split(
            "."
        )[0]

        if base_symbol.upper() not in lookup:

            lookup[
                base_symbol.upper()
            ] = symbol

    return lookup


# ==========================================================
# ΔΗΜΙΟΥΡΓΙΑ portfolio_updated.xlsx
# ==========================================================

def create_portfolio_excel(stocks):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Portfolio"

    # ------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------

    sheet.append(
        HEADERS
    )

    # ------------------------------------------------------
    # DATA
    # ------------------------------------------------------

    for stock in stocks:

        sheet.append([
            stock.get("symbol"),
            stock.get("name"),
            stock.get("sector"),
            stock.get("country"),
            stock.get("platform"),
            stock.get("purchase_date"),
            stock.get("purchase_price"),
            stock.get("quantity"),
            stock.get("current_price"),
            stock.get("market_value"),
            stock.get("percent_change"),
            stock.get("monthly_percent_change"),
            stock.get("portfolio_weight"),
            stock.get("vuaa_return"),
            stock.get("excess_return"),

            # Fundamentals
            stock.get("fcf_yield"),
            stock.get("fcf_growth"),
            stock.get("roic"),

            # Final Score
            stock.get("final_score"),
        ])

    # ------------------------------------------------------
    # FORMAT
    # ------------------------------------------------------

    format_portfolio_sheet(
        sheet
    )

    # ------------------------------------------------------
    # SAVE
    # ------------------------------------------------------

    workbook.save(
        OUTPUT_FILE
    )

    workbook.close()

    return OUTPUT_FILE


# ==========================================================
# ΕΝΗΜΕΡΩΣΗ portfolio.xlsx
# ==========================================================

def update_portfolio_excel(stocks):

    # ------------------------------------------------------
    # ΑΝ ΔΕΝ ΥΠΑΡΧΕΙ ΤΟ portfolio.xlsx
    # ------------------------------------------------------

    if not EXCEL_FILE.exists():

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Portfolio"

        sheet.append(
            HEADERS
        )

        workbook.save(
            EXCEL_FILE
        )

        workbook.close()

    # ------------------------------------------------------
    # ΑΝΟΙΓΟΥΜΕ ΤΟ ΥΠΑΡΧΟΝ EXCEL
    # ------------------------------------------------------

    workbook = load_workbook(
        EXCEL_FILE
    )

    sheet = workbook.worksheets[0]

    # ------------------------------------------------------
    # ΠΑΛΙΑ HEADERS
    #
    # Τα χρειαζόμαστε για να κρατήσουμε
    # χειροκίνητο sector / country.
    # ------------------------------------------------------

    old_headers = {}

    for cell in sheet[1]:

        if cell.value:

            header = str(
                cell.value
            ).strip().lower()

            old_headers[
                header
            ] = cell.column

    # ======================================================
    # ΚΡΑΤΑΜΕ MANUAL SECTOR / COUNTRY
    # ======================================================

    manual_data = {}

    if (
        "symbol" in old_headers
        and "platform" in old_headers
    ):

        for row in range(
            2,
            sheet.max_row + 1
        ):

            symbol = sheet.cell(
                row=row,
                column=old_headers["symbol"]
            ).value

            platform = sheet.cell(
                row=row,
                column=old_headers["platform"]
            ).value

            if not symbol:
                continue

            key = (
                str(
                    symbol
                ).strip().upper(),

                str(
                    platform or ""
                ).strip().upper()
            )

            sector = None
            country = None

            if "sector" in old_headers:

                sector = sheet.cell(
                    row=row,
                    column=old_headers["sector"]
                ).value

            if "country" in old_headers:

                country = sheet.cell(
                    row=row,
                    column=old_headers["country"]
                ).value

            manual_data[
                key
            ] = {
                "sector": sector,
                "country": country
            }

    # ======================================================
    # ΑΦΑΙΡΟΥΜΕ ΤΟ ΠΑΛΙΟ TABLE
    # ======================================================

    remove_existing_tables(
        sheet
    )

    # ======================================================
    # ΣΒΗΝΟΥΜΕ ΤΙΣ ΠΑΛΙΕΣ ΘΕΣΕΙΣ
    # ======================================================

    if sheet.max_row > 1:

        sheet.delete_rows(
            2,
            sheet.max_row - 1
        )

    # ======================================================
    # ΚΡΑΤΑΜΕ ΜΟΝΟ ΤΙΣ ΣΤΗΛΕΣ ΠΟΥ ΘΕΛΟΥΜΕ
    #
    # Πλέον:
    # A:S = 19 στήλες
    # ======================================================

    remove_extra_columns(
        sheet
    )

    # ======================================================
    # ΒΑΖΟΥΜΕ ΤΑ HEADERS ΣΤΗ ΣΩΣΤΗ ΣΕΙΡΑ
    # ======================================================

    write_correct_headers(
        sheet
    )

    # ======================================================
    # ΓΡΑΦΟΥΜΕ ΤΙΣ ΣΗΜΕΡΙΝΕΣ ΘΕΣΕΙΣ
    # ======================================================

    for stock in stocks:

        symbol = stock.get(
            "symbol"
        )

        platform = stock.get(
            "platform"
        )

        if not symbol:
            continue

        key = (
            str(
                symbol
            ).strip().upper(),

            str(
                platform or ""
            ).strip().upper()
        )

        # --------------------------------------------------
        # ΠΑΛΙΟ MANUAL SECTOR / COUNTRY
        # --------------------------------------------------

        old_manual_data = manual_data.get(
            key,
            {}
        )

        sector = old_manual_data.get(
            "sector"
        )

        country = old_manual_data.get(
            "country"
        )

        # --------------------------------------------------
        # ΝΕΟ SECTOR ΑΠΟ API/YAHOO
        # --------------------------------------------------

        if stock.get(
            "sector"
        ) is not None:

            sector = stock.get(
                "sector"
            )

        # --------------------------------------------------
        # ΝΕΟ COUNTRY ΑΠΟ API/YAHOO
        # --------------------------------------------------

        if stock.get(
            "country"
        ) is not None:

            country = stock.get(
                "country"
            )

        # --------------------------------------------------
        # ΓΡΑΦΗ ΣΤΟ EXCEL
        #
        # Η σειρά είναι ΑΚΡΙΒΩΣ ίδια με HEADERS.
        # --------------------------------------------------

        sheet.append([
            stock.get("symbol"),
            stock.get("name"),
            sector,
            country,
            stock.get("platform"),
            stock.get("purchase_date"),
            stock.get("purchase_price"),
            stock.get("quantity"),
            stock.get("current_price"),
            stock.get("market_value"),
            stock.get("percent_change"),
            stock.get("monthly_percent_change"),
            stock.get("portfolio_weight"),
            stock.get("vuaa_return"),
            stock.get("excess_return"),

            # Fundamentals
            stock.get("fcf_yield"),
            stock.get("fcf_growth"),
            stock.get("roic"),

            # Final Score
            stock.get("final_score"),
        ])

    # ======================================================
    # FORMAT + STOCKSUMMARY TABLE
    # ======================================================

    format_portfolio_sheet(
        sheet
    )

    # ======================================================
    # SAVE
    # ======================================================

    workbook.save(
        EXCEL_FILE
    )

    workbook.close()

    return EXCEL_FILE