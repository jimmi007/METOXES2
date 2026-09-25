from io import BytesIO
from metoxes.services.fundamentals_service import get_stock_fundamentals
from metoxes.services.analysis_service import analyze_portfolio
from fastapi import APIRouter, UploadFile, File, HTTPException
from metoxes.services.fundamentals_service import (
    get_stock_fundamentals,
    get_portfolio_fundamentals,
)
from openpyxl import load_workbook
from metoxes.services.portfolio_service import aggregate_positions
from metoxes.services.excel_service import (
    update_portfolio_excel,
    get_excel_stocks,
    create_portfolio_excel,
)
from metoxes.services.freedom_service import FREEDOM_TICKER_MAP
from metoxes.services.fundamentals_service import (
    get_stock_fundamentals,
    get_portfolio_fundamentals,
)

from metoxes.services.price_service import (
    get_monthly_percent_change,
)
from metoxes.services.scoring_services import score_stocks_by_sector

from metoxes.services.trading212_service import (
    get_positions,
    get_clean_positions,
)

from metoxes.services.capital_service import (
    create_capital_session,
    get_capital_positions,
    get_clean_capital_positions,
)

from metoxes.services.freedom_service import (
    test_freedom_connection,
    get_freedom_positions,
    get_freedom_trades,
    get_clean_freedom_positions,
    get_freedom_connection,
)


router = APIRouter()


# ==========================================================
# BASIC TEST
# ==========================================================





@router.get("/stocks")
async def get_stocks():

    return {
        "message": "METOXES2 REST API works"
    }


# ==========================================================
# TRADING212 - RAW POSITIONS
# ==========================================================

@router.get("/stocks/trading212")
async def get_trading212_positions():

    positions = await get_positions()

    return {
        "count": len(positions),
        "positions": positions
    }


# ==========================================================
# READ EXCEL
# ==========================================================

@router.post("/stocks/read-excel")
async def read_excel(
    file: UploadFile = File(...)
):

    if not file.filename.endswith(
        (".xlsx", ".xlsm")
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Το αρχείο πρέπει να είναι "
                "Excel (.xlsx ή .xlsm)"
            )
        )

    contents = await file.read()

    workbook = load_workbook(
        BytesIO(contents),
        data_only=True
    )

    worksheet = workbook.active

    headers = []

    for cell in worksheet[1]:

        if cell.value is not None:
            headers.append(
                str(cell.value)
                .strip()
                .lower()
            )

        else:
            headers.append(None)

    required_columns = [
        "symbol",
        "name",
        "sector",
        "country",
        "platform",
        "purchase_date",
        "purchase_price",
        "quantity",
        "current_price",
        "market_value",
        "percent_change",
        "portfolio_weight",
        "vuaa_return",
        "excess_return",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in headers
    ]

    if missing_columns:

        raise HTTPException(
            status_code=400,
            detail={
                "message":
                    "Λείπουν στήλες από το Excel",

                "missing_columns":
                    missing_columns
            }
        )

    stocks = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not any(row):
            continue

        stock = {}

        for index, header in enumerate(headers):

            if header is not None:
                stock[header] = row[index]

        if not stock.get("symbol"):
            continue

        stocks.append(stock)

    return {
        "filename": file.filename,
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# MONTHLY CHANGE TEST
# ==========================================================

@router.get("/stocks/monthly/{symbol}")
async def get_monthly_change(
    symbol: str
):

    monthly_change = (
        get_monthly_percent_change(symbol)
    )

    return {
        "symbol": symbol,
        "monthly_percent_change":
            monthly_change
    }


# ==========================================================
# EXCEL SYMBOLS
# ==========================================================

@router.get("/stocks/excel-symbols")
async def get_symbols_from_excel():

    stocks = get_excel_stocks()

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# TRADING212 - CLEAN
# ==========================================================

@router.get("/stocks/clean")
async def get_clean_stocks():

    stocks = await get_clean_positions()

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# CREATE EXCEL
# ==========================================================

@router.post("/stocks/create-excel")
async def create_excel():

    stocks = await get_clean_positions()

    file_path = create_portfolio_excel(
        stocks
    )

    return {
        "message":
            "Excel created successfully",

        "stocks":
            len(stocks),

        "file":
            str(file_path)
    }


# ==========================================================
# UPDATE EXCEL
# ==========================================================

# ==========================================================
# UPDATE EXCEL
# ==========================================================

# ==========================================================
# UPDATE EXCEL - ΟΛΟ ΤΟ PORTFOLIO
# ==========================================================

@router.post("/stocks/update-excel")
async def update_excel():

    # Παίρνουμε Trading212
    trading212_stocks = await get_clean_positions()

    # Παίρνουμε Capital
    capital_stocks = await get_clean_capital_positions()

    # Παίρνουμε Freedom24
    freedom_stocks = get_clean_freedom_positions()

    # Ενώνουμε όλες τις θέσεις
    all_stocks = (
        trading212_stocks
        + capital_stocks
        + freedom_stocks
    )

    # Ενώνουμε πολλαπλές θέσεις της ίδιας
    # μετοχής στον ίδιο broker
    all_stocks = aggregate_positions(
        all_stocks
    )

    # Συνολική αξία portfolio σε EUR
    total_portfolio_value = sum(
        stock["market_value"]
        for stock in all_stocks
        if stock.get("market_value") is not None
    )

    # Υπολογίζουμε portfolio weight
    for stock in all_stocks:

        market_value = stock.get(
            "market_value"
        )

        if (
            market_value is not None
            and total_portfolio_value > 0
        ):

            stock["portfolio_weight"] = round(
                market_value
                / total_portfolio_value
                * 100,
                2
            )

        else:

            stock["portfolio_weight"] = None

    # Γράφουμε όλα τα δεδομένα στο Excel
    file_path = update_portfolio_excel(
        all_stocks
    )

    # Απάντηση API
    return {
        "message":
            "Portfolio Excel updated successfully",

        "stocks":
            len(all_stocks),

        "total_portfolio_value":
            round(
                total_portfolio_value,
                2
            ),

        "file":
            str(file_path)
    }

# ==========================================================
# CAPITAL - TEST
# ==========================================================

@router.get("/stocks/capital/test")
async def test_capital():

    result = await create_capital_session()

    return result


# ==========================================================
# CAPITAL - RAW POSITIONS
# ==========================================================

@router.get("/stocks/capital")
async def capital_positions():

    positions = await get_capital_positions()

    return positions


# ==========================================================
# CAPITAL - CLEAN
# ==========================================================

@router.get("/stocks/capital/clean")
async def clean_capital_positions():

    stocks = (
        await get_clean_capital_positions()
    )

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# FREEDOM24 - TEST
# ==========================================================

@router.get("/stocks/freedom/test")
async def test_freedom():

    return test_freedom_connection()


# ==========================================================
# FREEDOM24 - RAW POSITIONS
# ==========================================================

@router.get("/stocks/freedom")
async def freedom_positions():

    positions = get_freedom_positions()

    return {
        "count": len(positions),
        "positions": positions
    }


# ==========================================================
# FREEDOM24 - TRADES
# ==========================================================

@router.get("/stocks/freedom/trades")
async def freedom_trades():

    trades = get_freedom_trades()

    return {
        "count": len(trades),
        "trades": trades
    }


# ==========================================================
# FREEDOM24 - CLEAN
# ==========================================================

@router.get("/stocks/freedom/clean")
async def clean_freedom_positions():

    stocks = get_clean_freedom_positions()

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# FREEDOM24 - DEBUG CURRENCIES
# ==========================================================

@router.get("/stocks/freedom/debug")
async def freedom_debug():

    connection = get_freedom_connection()

    response = (
        connection.authorized_request(
            "getPositionJson"
        )
    )

    portfolio = response["result"]["ps"]

    accounts = portfolio.get(
        "acc",
        []
    )

    return {
        "accounts_count":
            len(accounts),

        "account_keys": (
            list(accounts[0].keys())
            if accounts
            else []
        ),

        "currencies": [
            {
                "curr":
                    account.get("curr"),

                "currval":
                    account.get("currval")
            }
            for account in accounts
        ]
    }


# ==========================================================
# FREEDOM24 - TRADES DEBUG
# ==========================================================

@router.get("/stocks/freedom/trades-debug")
async def freedom_trades_debug():

    connection = get_freedom_connection()

    response = (
        connection.authorized_request(
            "getTradesHistory",
            {
                "beginDate":
                    "2020-01-01T00:00:00",

                "endDate":
                    "2030-12-31T23:59:59",

                "max":
                    100,

                "sort":
                    1,
            }
        )
    )

    trades_section = response.get(
        "trades",
        {}
    )

    trades = trades_section.get(
        "trade",
        []
    )

    return {
        "max_trade_id":
            response.get("max_trade_id"),

        "trade_count": (
            len(trades)
            if isinstance(trades, list)
            else 1
        ),

        "first_trade_id": (
            trades[0].get("id")
            if (
                isinstance(trades, list)
                and trades
            )
            else None
        ),

        "last_trade_id": (
            trades[-1].get("id")
            if (
                isinstance(trades, list)
                and trades
            )
            else None
        )
    }
@router.get("/stocks/all")
async def get_all_stocks():

    # Trading212
    trading212_stocks = await get_clean_positions()

    # Capital
    capital_stocks = await get_clean_capital_positions()

    # Freedom24
    freedom_stocks = get_clean_freedom_positions()

    # Όλες οι θέσεις μαζί
    all_stocks = (
        trading212_stocks
        + capital_stocks
        + freedom_stocks
    )
    # Ενώνουμε τις πολλαπλές θέσεις
    # της ίδιας μετοχής στον ίδιο broker
    all_stocks = aggregate_positions(
        all_stocks
    )
    # ==================================================
    # ΣΥΝΟΛΙΚΗ ΑΞΙΑ ΧΑΡΤΟΦΥΛΑΚΙΟΥ
    # Όλα τα market_value είναι πλέον σε EUR
    # ==================================================

    total_portfolio_value = sum(
        stock["market_value"]
        for stock in all_stocks
        if stock.get("market_value") is not None
    )

    # ==================================================
    # PORTFOLIO WEIGHT
    # Βάρος κάθε θέσης στο συνολικό χαρτοφυλάκιο
    # ==================================================

    for stock in all_stocks:

        market_value = stock.get("market_value")

        if (
                market_value is not None
                and total_portfolio_value > 0
        ):

            stock["portfolio_weight"] = round(
                market_value
                / total_portfolio_value
                * 100,
                2
            )

        else:

            stock["portfolio_weight"] = None
    return {
        "count": len(all_stocks),
        "total_portfolio_value": round(
            total_portfolio_value,
            2
        ),
        "stocks": all_stocks
    }


@router.get("/stocks/analysis")
async def portfolio_analysis():

    trading212 = await get_clean_positions()
    capital = await get_clean_capital_positions()
    freedom = get_clean_freedom_positions()

    all_stocks = (
        trading212
        + capital
        + freedom
    )

    aggregated = aggregate_positions(
        all_stocks
    )

    return analyze_portfolio(
        aggregated
    )


@router.get("/stocks/fundamentals/{symbol}")
async def stock_fundamentals(symbol: str):

    return get_stock_fundamentals(
        symbol.upper()
    )

@router.get("/stocks/fundamentals")
def get_portfolio_fundamentals(stocks):

    results = []

    for stock in stocks:

        symbol = stock.get("symbol")
        platform = stock.get("platform")

        if not symbol:
            continue

        # --------------------------------------------------
        # YAHOO SYMBOL
        # --------------------------------------------------

        yahoo_symbol = symbol

        # Freedom24 χρησιμοποιεί δικά της symbols:
        # ACM.US -> ACM
        # ASML.EU -> ASML
        # AEGN.GR -> AEGN.AT
        # κτλ.
        if platform == "Freedom24":

            yahoo_symbol = FREEDOM_TICKER_MAP.get(
                symbol,
                symbol
            )

        try:

            fundamentals = get_stock_fundamentals(
                yahoo_symbol
            )

            results.append({

                # Κρατάμε το αρχικό broker symbol
                "symbol": symbol,

                # Για έλεγχο βλέπουμε και τι στείλαμε Yahoo
                "yahoo_symbol": yahoo_symbol,

                "name": stock.get("name"),
                "sector": stock.get("sector"),
                "country": stock.get("country"),
                "platform": platform,

                "fcf_yield": fundamentals.get(
                    "fcf_yield"
                ),

                "fcf_growth": fundamentals.get(
                    "fcf_growth"
                ),

                "roic": fundamentals.get(
                    "roic"
                ),
            })

        except Exception as e:

            print(
                f"FUNDAMENTALS ERROR "
                f"{symbol} -> {yahoo_symbol}: {e}"
            )

            results.append({

                "symbol": symbol,
                "yahoo_symbol": yahoo_symbol,
                "name": stock.get("name"),
                "sector": stock.get("sector"),
                "country": stock.get("country"),
                "platform": platform,

                "fcf_yield": None,
                "fcf_growth": None,
                "roic": None,
            })

    return results
@router.get("/stocks/scores")
async def portfolio_scores():

    # ----------------------------------------------
    # ΠΑΙΡΝΟΥΜΕ ΤΙΣ ΘΕΣΕΙΣ ΑΠΟ ΤΙΣ 3 ΠΛΑΤΦΟΡΜΕΣ
    # ----------------------------------------------

    trading212 = await get_clean_positions()
    capital = await get_clean_capital_positions()
    freedom = get_clean_freedom_positions()

    all_stocks = (
        trading212
        + capital
        + freedom
    )

    # ----------------------------------------------
    # ΟΜΑΔΟΠΟΙΗΣΗ ΘΕΣΕΩΝ
    # ----------------------------------------------

    aggregated = aggregate_positions(
        all_stocks
    )

    # ----------------------------------------------
    # FUNDAMENTALS
    # ----------------------------------------------

    fundamentals = get_portfolio_fundamentals(
        aggregated
    )

    # ----------------------------------------------
    # SCORE ΑΝΑ SECTOR
    # ----------------------------------------------

    scores = score_stocks_by_sector(
        fundamentals
    )

    return {
        "count": len(scores),
        "stocks": scores
    }