import asyncio
import json
import os
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from metoxes.services.fundamentals_service import get_stock_fundamentals
from metoxes.services.analysis_service import analyze_portfolio
from fastapi import APIRouter, UploadFile, File, HTTPException
from metoxes.services.fundamentals_service import (
    get_stock_fundamentals,
    get_portfolio_fundamentals,
)
from metoxes.services.data_quality_service import (
    run_portfolio_quality_checks,
    save_quality_report,
)

from metoxes.services.dashboard_service import (
    generate_portfolio_dashboard,
    refresh_dashboard_candidates,
)
from metoxes.services.forward_growth_service import (
    add_forward_growth_scores,
)
from metoxes.services.absolute_quality_service import (
    add_absolute_quality_scores,
)
from metoxes.services.portfolio_metadata_service import (
    enrich_sector_country_from_saved_positions,
)
from metoxes.services.candidate_research_service import (
    research_top_candidates,
)
from metoxes.services.sold_history_service import (
    safe_collect_freedom_realized_history,
    safe_collect_capital_realized_history,
    safe_collect_trading212_realized_history,
    merge_realized_history_results,
)
from metoxes.services.broker_history_debug_service import (
    safe_fetch_trading212_history_debug,
    safe_fetch_capital_transactions_debug,
)
from metoxes.services.secondary_fundamentals_service import (
    enrich_with_secondary_fundamentals,
)
from metoxes.services.financial_sector_scoring_service import (
    apply_financial_sector_model,
)
from metoxes.services.anomaly_detection_service import (
    detect_portfolio_anomalies,
)
from metoxes.services.market_context_service import (
    enrich_market_context,
)
from metoxes.services.ai_commentary_service import (
    enrich_ai_commentary,
)
from metoxes.services.peer_intelligence_service import (
    enrich_peer_intelligence,
)
from metoxes.services.event_impact_service import (
    enrich_event_impact,
)
from metoxes.services.scenario_analysis_service import (
    enrich_standard_scenarios,
    analyze_custom_scenario,
)
from metoxes.services.model_backtest_service import (
    record_model_snapshots,
    run_score_backtest,
    load_latest_backtest_report,
)
from metoxes.services.score_explanation_service import (
    enrich_score_explanations,
    reconcile_financial_crosschecks,
)
from metoxes.services.portfolio_service import aggregate_positions
from metoxes.services.excel_service import (
    update_portfolio_excel,
    update_candidate_research_excel,
    get_excel_stocks,
    create_portfolio_excel,
)
from metoxes.services.freedom_service import FREEDOM_TICKER_MAP
from metoxes.services.fundamentals_service import (
    get_stock_fundamentals,
    get_portfolio_fundamentals,
)
from metoxes.services.fundamentals_service import (
    get_stock_fundamentals,
    get_portfolio_fundamentals,
)


from metoxes.services.price_service import (
    get_monthly_percent_change,
)
from metoxes.services.scoring_services import score_stocks_by_sector

from metoxes.services.trading212_service import (
    get_positions,
    get_clean_positions,
)

from metoxes.services.capital_service import (
    create_capital_session,
    get_capital_positions,
    get_clean_capital_positions,
)

from metoxes.services.freedom_service import (
    test_freedom_connection,
    get_freedom_positions,
    get_freedom_trades,
    get_clean_freedom_positions,
    get_freedom_connection,
)


router = APIRouter()

BUILD_VERSION = "V21-quarter-beta-investing"


# ==========================================================
# ANALYSIS ENRICHMENT STATUS
# ==========================================================

@router.get("/stocks/enrichment-status")
async def enrichment_status():
    """Visible health/status endpoint for the current METOXES2 build."""
    return {
        "status": "OK",
        "build": BUILD_VERSION,
        "expected_routes": [
            "/stocks/enrichment-status",
            "/stocks/model-intelligence",
            "/stocks/scenario/{symbol}",
            "/stocks/backtest-scores",
            "/stocks/capital/realized",
            "/stocks/freedom/realized",
            "/stocks/trading212/history-debug",
            "/stocks/trading212/realized",
            "/stocks/update-excel",
        ],
        "openai_configured": bool(os.getenv("OPENAI_API_KEY")),
        "openai_model": os.getenv("OPENAI_MODEL", "gpt-6-luna"),
        "fmp_configured": bool(os.getenv("FMP_API_KEY")),
        "fmp_news_enabled": os.getenv("FMP_NEWS_ENABLED", "0").lower() in {"1", "true", "yes"},
        "sec_user_agent_configured": bool(os.getenv("SEC_USER_AGENT")),
        "sold_history": {
            "freedom24": "enabled",
            "capital": "enabled",
            "trading212": "enabled from FILLED SELL fill.walletImpact.realisedProfitLoss",
        },
        "features": [
            "AI commentary per stock",
            "real competitor / peer intelligence",
            "AI + deterministic news and earnings impact scoring",
            "scenario stress tests with shadow scores",
            "point-in-time score history and no-look-ahead backtesting",
            "predictive metric IC and weight recommendations",
            "news and earnings context",
            "dedicated Financial Services scoring",
            "anomaly detection",
            "independent fundamentals cross-check",
            "Capital realized P/L from Trade closed transactions",
            "Freedom24 realized P/L from getTradesHistory",
            "Trading212 realized P/L from walletImpact.realisedProfitLoss",
            "Master collapsed accordions for existing and candidate stocks",
            "Partial sales shown in Sold while remaining quantity stays in Portfolio",
            "Sold KPI formulas from quantity x purchase/sale price",
            "Detailed candidate score/news/FMP/SEC/anomaly evaluation",
        ],
    }


def enrich_stocks_with_saved_metadata(stocks):
    """
    Broker metadata έχει πάντα προτεραιότητα.

    Αν όμως sector/country λείπουν (κυρίως Freedom24),
    τα παίρνουμε από το υπάρχον portfolio.xlsx πριν
    αρχίσει το relative scoring.
    """
    try:
        saved_stocks = (
            get_excel_stocks()
        )
    except (
        FileNotFoundError,
        ValueError,
    ):
        return stocks
    except Exception as error:
        print(
            "PORTFOLIO METADATA ERROR: "
            f"{error}"
        )
        return stocks

    return (
        enrich_sector_country_from_saved_positions(
            stocks,
            saved_stocks,
        )
    )


# ==========================================================
# BASIC TEST
# ==========================================================





@router.get("/stocks")
async def get_stocks():

    return {
        "message": "METOXES2 REST API works"
    }


# ==========================================================
# TRADING212 - RAW POSITIONS
# ==========================================================

@router.get("/stocks/trading212")
async def get_trading212_positions():

    positions = await get_positions()

    return {
        "count": len(positions),
        "positions": positions
    }


# ==========================================================
# READ EXCEL
# ==========================================================

@router.post("/stocks/read-excel")
async def read_excel(
    file: UploadFile = File(...)
):

    if not file.filename.endswith(
        (".xlsx", ".xlsm")
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Το αρχείο πρέπει να είναι "
                "Excel (.xlsx ή .xlsm)"
            )
        )

    contents = await file.read()

    workbook = load_workbook(
        BytesIO(contents),
        data_only=True
    )

    worksheet = workbook.active

    headers = []

    for cell in worksheet[1]:

        if cell.value is not None:
            headers.append(
                str(cell.value)
                .strip()
                .lower()
            )

        else:
            headers.append(None)

    required_columns = [
        "symbol",
        "name",
        "sector",
        "country",
        "platform",
        "purchase_date",
        "purchase_price",
        "quantity",
        "current_price",
        "market_value",
        "percent_change",
        "portfolio_weight",
        "vuaa_return",
        "excess_return",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in headers
    ]

    if missing_columns:

        raise HTTPException(
            status_code=400,
            detail={
                "message":
                    "Λείπουν στήλες από το Excel",

                "missing_columns":
                    missing_columns
            }
        )

    stocks = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not any(row):
            continue

        stock = {}

        for index, header in enumerate(headers):

            if header is not None:
                stock[header] = row[index]

        if not stock.get("symbol"):
            continue

        stocks.append(stock)

    return {
        "filename": file.filename,
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# MONTHLY CHANGE TEST
# ==========================================================

@router.get("/stocks/monthly/{symbol}")
async def get_monthly_change(
    symbol: str
):

    monthly_change = (
        get_monthly_percent_change(symbol)
    )

    return {
        "symbol": symbol,
        "monthly_percent_change":
            monthly_change
    }


# ==========================================================
# EXCEL SYMBOLS
# ==========================================================

@router.get("/stocks/excel-symbols")
async def get_symbols_from_excel():

    stocks = get_excel_stocks()

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# TRADING212 - CLEAN
# ==========================================================

@router.get("/stocks/clean")
async def get_clean_stocks():

    stocks = await get_clean_positions()

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# CREATE EXCEL
# ==========================================================

@router.post("/stocks/create-excel")
async def create_excel():

    stocks = await get_clean_positions()

    file_path = create_portfolio_excel(
        stocks
    )

    return {
        "message":
            "Excel created successfully",

        "stocks":
            len(stocks),

        "file":
            str(file_path)
    }


# ==========================================================
# UPDATE EXCEL
# ==========================================================

# ==========================================================
# UPDATE EXCEL
# ==========================================================

# ==========================================================
# UPDATE EXCEL - ΟΛΟ ΤΟ PORTFOLIO
# ==========================================================

@router.post("/stocks/update-excel")
async def update_excel():

    # ======================================================
    # 1. TRADING212
    # ======================================================

    trading212 = (
        await get_clean_positions()
    )

    # ======================================================
    # 2. CAPITAL
    # ======================================================

    capital = (
        await get_clean_capital_positions()
    )

    # ======================================================
    # 3. FREEDOM24
    # ======================================================

    freedom = (
        get_clean_freedom_positions()
    )

    # ======================================================
    # 4. ΟΛΕΣ ΟΙ ΘΕΣΕΙΣ
    # ======================================================

    all_stocks = (
        trading212
        + capital
        + freedom
    )

    # ======================================================
    # 5. AGGREGATION
    #
    # ίδιο symbol + ίδιο platform
    # ======================================================

    stocks = aggregate_positions(
        all_stocks
    )

    # ======================================================
    # 5A. SECTOR / COUNTRY FALLBACK ΑΠΟ portfolio.xlsx
    # ======================================================
    #
    # Η Freedom24 δεν επιστρέφει πάντα sector/country.
    # Τα συμπληρώνουμε ΠΡΙΝ από fundamentals/scoring,
    # ώστε το Relative Score να χρησιμοποιεί σωστό
    # country + sector peer group.
    # ======================================================

    stocks = enrich_stocks_with_saved_metadata(
        stocks
    )

    # ======================================================
    # 6. TOTAL PORTFOLIO VALUE
    # ======================================================

    total_portfolio_value = sum(
        stock.get(
            "market_value"
        ) or 0
        for stock in stocks
    )

    # ======================================================
    # 7. PORTFOLIO WEIGHT
    # ======================================================

    for stock in stocks:

        market_value = stock.get(
            "market_value"
        )

        if (
            market_value is not None
            and total_portfolio_value > 0
        ):

            stock[
                "portfolio_weight"
            ] = round(
                market_value
                / total_portfolio_value
                * 100,
                2
            )

        else:

            stock[
                "portfolio_weight"
            ] = None

    # ======================================================
    # 8. FUNDAMENTALS
    # ======================================================

    fundamentals = (
        get_portfolio_fundamentals(
            stocks
        )
    )

    # ======================================================
    # 9. SECTOR SCORES
    # ======================================================

    scores = (
        score_stocks_by_sector(
            fundamentals
        )
    )

    scores = add_forward_growth_scores(
        scores
    )

    scores = add_absolute_quality_scores(
        scores
    )

    # ======================================================
    # 9A. INDEPENDENT FUNDAMENTALS CROSS-CHECK
    # ======================================================
    # FMP when configured; SEC EDGAR fallback for US issuers.
    scores = await asyncio.to_thread(
        enrich_with_secondary_fundamentals,
        scores,
    )

    # ======================================================
    # 9B. BANKS / FINANCIAL SERVICES SPECIAL MODEL
    # ======================================================
    # Overrides the generic FCF/ROIC score only for Financial Services.
    scores = await asyncio.to_thread(
        apply_financial_sector_model,
        scores,
    )

    # Financial Services cross-check must use ROE/P-B, not generic FCF/ROIC.
    scores = reconcile_financial_crosschecks(
        scores
    )

    # ======================================================
    # 9C. PUBLIC SCORE EXPLANATION
    # ======================================================
    # Creates human-readable drivers/risks without exposing
    # hidden Relative / Absolute score numbers.
    scores = enrich_score_explanations(
        scores
    )

    # ======================================================
    # 10. SCORE LOOKUP
    #
    # symbol + platform
    # ======================================================

    score_lookup = {}

    for score in scores:

        key = (
            str(
                score.get(
                    "symbol"
                ) or ""
            ).strip().upper(),

            str(
                score.get(
                    "platform"
                ) or ""
            ).strip().upper()
        )

        score_lookup[
            key
        ] = score

    # ======================================================
    # 11. ΠΡΟΣΘΗΚΗ FUNDAMENTALS ΣΤΙΣ ΘΕΣΕΙΣ
    # ======================================================

    for stock in stocks:

        key = (
            str(
                stock.get(
                    "symbol"
                ) or ""
            ).strip().upper(),

            str(
                stock.get(
                    "platform"
                ) or ""
            ).strip().upper()
        )

        score = score_lookup.get(
            key,
            {}
        )

        stock[
            "fcf_yield"
        ] = score.get(
            "fcf_yield"
        )

        stock[
            "fcf_growth"
        ] = score.get(
            "fcf_growth"
        )

        stock[
            "roic"
        ] = score.get(
            "roic"
        )

        stock[
            "forward_revenue_growth"
        ] = score.get(
            "forward_revenue_growth"
        )

        stock[
            "forward_eps_growth"
        ] = score.get(
            "forward_eps_growth"
        )

        stock[
            "relative_score"
        ] = score.get(
            "relative_score"
        )

        stock[
            "absolute_score"
        ] = score.get(
            "absolute_score"
        )

        stock[
            "absolute_coverage"
        ] = score.get(
            "absolute_coverage"
        )

        stock[
            "score_mode"
        ] = score.get(
            "score_mode"
        )

        stock[
            "final_score"
        ] = score.get(
            "final_score"
        )

        # Extra diagnostics stay out of the visible Excel score columns
        # but feed anomaly detection and AI commentary.
        for analysis_field in (
            "secondary_source",
            "secondary_status",
            "secondary_fcf_yield",
            "secondary_roic",
            "secondary_roe",
            "secondary_price_to_book",
            "secondary_fcf_ttm",
            "secondary_revenue_ttm",
            "secondary_net_income_ttm",
            "fundamental_crosscheck_status",
            "fundamental_discrepancy_pct",
            "financial_model",
            "financial_industry",
            "financial_roe",
            "financial_price_to_book",
            "financial_profit_margin",
            "financial_revenue_growth",
            "financial_forward_eps_growth",
            "financial_model_coverage",
            "financial_score_mode",
            "score_rationale_headline",
            "score_rationale_text",
            "score_drivers",
            "score_weaknesses",
            "score_missing_metrics",
            "score_metric_details",
            "score_crosscheck_text",
            "score_model_label",
        ):
            stock[analysis_field] = score.get(analysis_field)

    # ======================================================
    # 11A. ANOMALY DETECTION
    # ======================================================
    anomaly_result = detect_portfolio_anomalies(
        stocks
    )
    stocks = anomaly_result["stocks"]

    # ======================================================
    # 11B. NEWS + EARNINGS CONTEXT
    # ======================================================
    stocks = await asyncio.to_thread(
        enrich_market_context,
        stocks,
    )

    # ======================================================
    # 11C. REAL COMPETITOR / PEER INTELLIGENCE
    # ======================================================
    stocks = await asyncio.to_thread(
        enrich_peer_intelligence,
        stocks,
        4,
        20,
        5,
    )

    # ======================================================
    # 11D. NEWS / EARNINGS IMPACT
    # ======================================================
    # Separate advisory layer (-100..+100). It does not overwrite Final Score.
    stocks = await asyncio.to_thread(
        enrich_event_impact,
        stocks,
    )

    # ======================================================
    # 11E. STANDARD SCENARIO STRESS TESTS
    # ======================================================
    stocks = enrich_standard_scenarios(
        stocks
    )

    # ======================================================
    # 11F. AI COMMENTARY PER STOCK
    # ======================================================
    stocks = await asyncio.to_thread(
        enrich_ai_commentary,
        stocks,
    )

    # ======================================================
    # 11G. POINT-IN-TIME MODEL HISTORY
    # ======================================================
    # Needed for a genuine forward-return backtest without look-ahead bias.
    try:
        model_snapshot_result = record_model_snapshots(
            stocks,
            source="portfolio",
        )
    except Exception as error:
        model_snapshot_result = {
            "status": "error",
            "error": str(error)[:300],
        }

    # ======================================================
    # 12. DATA QUALITY CHECKS
    # ======================================================

    quality_report = (
        run_portfolio_quality_checks(
            stocks
        )
    )

    quality_report["anomaly_detection"] = (
        anomaly_result.get("report", {})
    )

    quality_file = (
        save_quality_report(
            quality_report
        )
    )

    # ======================================================
    # 13. REALIZED SOLD HISTORY
    # ======================================================
    # All three brokers are now broker-history backed.
    # Trading212 uses only FILLED SELL fills and reads the broker-reported
    # fill.walletImpact.realisedProfitLoss value. All cursor pages are fetched.
    freedom_history_task = asyncio.to_thread(
        safe_collect_freedom_realized_history
    )
    capital_history_task = safe_collect_capital_realized_history()
    trading212_history_task = safe_collect_trading212_realized_history()

    freedom_history, capital_history, trading212_history = await asyncio.gather(
        freedom_history_task,
        capital_history_task,
        trading212_history_task,
    )

    sold_history_result = merge_realized_history_results(
        freedom_history,
        capital_history,
        trading212_history,
    )

    # ======================================================
    # 14. UPDATE PORTFOLIO EXCEL
    # ======================================================

    file_path = (
        update_portfolio_excel(
            stocks,
            sold_realized_lookup=sold_history_result.get(
                "lookup",
                {},
            ),
        )
    )

    # ======================================================
    # 15. CREATE JSON + HTML DASHBOARD
    # ======================================================

    dashboard_files = (
        generate_portfolio_dashboard(
            stocks=stocks,
            quality_report=quality_report,
        )
    )

    # ======================================================
    # 16. RESPONSE
    # ======================================================

    return {

        "build": BUILD_VERSION,

        "message":
            "Portfolio Excel + dashboard updated successfully",

        "stocks":
            len(stocks),

        "total_portfolio_value":
            round(
                total_portfolio_value,
                2
            ),

        "sold_history":
            sold_history_result.get(
                "meta",
                {}
            ),

        "analysis_enrichment": {
            "openai_commentary": bool(
                os.getenv("OPENAI_API_KEY")
            ),
            "openai_model": os.getenv(
                "OPENAI_MODEL",
                "gpt-6-luna",
            ),
            "fmp_configured": bool(
                os.getenv("FMP_API_KEY")
            ),
            "secondary_fundamentals": (
                "FMP or SEC EDGAR"
            ),
            "anomaly_detection": anomaly_result.get(
                "report",
                {},
            ),
            "peer_intelligence": True,
            "event_impact": True,
            "scenario_analysis": True,
            "model_history": model_snapshot_result,
            "latest_backtest": load_latest_backtest_report(),
        },

        "data_quality": {

            "status":
                quality_report[
                    "status"
                ],

            "errors":
                quality_report[
                    "errors"
                ],

            "warnings":
                quality_report[
                    "warnings"
                ],

            "info":
                quality_report[
                    "info"
                ],

            "positions_with_issues":
                quality_report[
                    "positions_with_issues"
                ],
        },

        "files": {

            "excel":
                str(
                    file_path
                ),

            "quality_report":
                str(
                    quality_file
                ),

            "dashboard_json":
                str(
                    dashboard_files[
                        "json_file"
                    ]
                ),

            "dashboard_html":
                str(
                    dashboard_files[
                        "html_file"
                    ]
                ),
        }
    }


# ==========================================================
# V20 MODEL INTELLIGENCE / SCENARIO / BACKTEST
# ==========================================================

@router.get("/stocks/model-intelligence")
async def model_intelligence_status():
    return {
        "build": BUILD_VERSION,
        "status": "OK",
        "live_score_policy": "deterministic Final Score unchanged by AI/event layer",
        "peer_intelligence": "FMP peers when configured; analysed-universe fallback",
        "event_impact": "-100..+100 advisory impact with OpenAI/rules fallback",
        "scenario_analysis": "shadow scores; does not overwrite Final Score",
        "backtesting": "point-in-time snapshots -> forward returns -> Spearman IC -> suggested weights",
        "latest_backtest": load_latest_backtest_report(),
    }


@router.get("/stocks/scenario/{symbol}")
async def stock_scenario(
    symbol: str,
    eps_growth_shock_pp: float = -15.0,
    revenue_growth_shock_pp: float = 0.0,
    fcf_growth_shock_pp: float = 0.0,
    roic_change_pct: float = 0.0,
    roe_change_pct: float = 0.0,
    profit_margin_shock_pp: float = 0.0,
):
    """Run a custom shadow scenario from the latest dashboard snapshot."""
    dashboard_file = Path(__file__).resolve().parents[2] / "portfolio_dashboard_data.json"
    if not dashboard_file.exists():
        raise HTTPException(status_code=404, detail="Run POST /stocks/update-excel first.")
    try:
        payload = json.loads(dashboard_file.read_text(encoding="utf-8"))
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"Cannot read dashboard data: {error}")

    wanted = str(symbol or "").strip().upper()
    row = None
    source = None
    for bucket in ("stocks", "candidates"):
        for item in payload.get(bucket) or []:
            if str(item.get("symbol") or "").strip().upper() == wanted:
                row = item
                source = bucket
                break
        if row is not None:
            break
    if row is None:
        raise HTTPException(status_code=404, detail=f"Symbol {wanted} not found in latest dashboard data.")

    result = analyze_custom_scenario(
        row,
        eps_growth_shock_pp=eps_growth_shock_pp,
        revenue_growth_shock_pp=revenue_growth_shock_pp,
        fcf_growth_shock_pp=fcf_growth_shock_pp,
        roic_change_pct=roic_change_pct,
        roe_change_pct=roe_change_pct,
        profit_margin_shock_pp=profit_margin_shock_pp,
        name="API custom scenario",
    )
    return {
        "build": BUILD_VERSION,
        "symbol": wanted,
        "source": source,
        "result": result,
        "note": "Scenario score is a shadow analysis and does not change Final Score.",
    }


@router.post("/stocks/backtest-scores")
async def backtest_scores(
    horizon_days: int = 90,
    min_samples: int = 12,
    benchmark: str | None = None,
):
    if horizon_days < 20 or horizon_days > 730:
        raise HTTPException(status_code=400, detail="horizon_days must be between 20 and 730.")
    if min_samples < 6 or min_samples > 1000:
        raise HTTPException(status_code=400, detail="min_samples must be between 6 and 1000.")
    try:
        report = await asyncio.to_thread(
            run_score_backtest,
            horizon_days,
            min_samples,
            benchmark,
        )
    except Exception as error:
        raise HTTPException(status_code=503, detail=f"Backtest failed: {error}")
    # Refresh the existing dashboard immediately so the Model Intelligence
    # accordion shows the new backtest report without requiring another
    # broker refresh.
    try:
        dashboard_file = Path(__file__).resolve().parents[2] / "portfolio_dashboard_data.json"
        if dashboard_file.exists():
            current = json.loads(dashboard_file.read_text(encoding="utf-8"))
            generate_portfolio_dashboard(
                stocks=current.get("stocks", []),
                quality_report=current.get("quality", {}),
                candidates=current.get("candidates", []),
            )
    except Exception:
        pass

    return {
        "build": BUILD_VERSION,
        "report": report,
    }


# ==========================================================
# FREEDOM24 - REALIZED SOLD HISTORY
# ==========================================================

@router.get("/stocks/freedom/realized")
async def freedom_realized_history():
    result = await asyncio.to_thread(
        safe_collect_freedom_realized_history
    )

    return {
        "meta": result.get("meta", {}),
        "count": len(result.get("summaries", [])),
        "stocks": result.get("summaries", []),
    }


# ==========================================================
# CAPITAL - VERIFIED REALIZED P/L
# ==========================================================

@router.get("/stocks/capital/realized")
async def capital_realized_history():
    result = await safe_collect_capital_realized_history()
    return {
        "build": BUILD_VERSION,
        "meta": result.get("meta", {}),
        "count": len(result.get("summaries", [])),
        "stocks": result.get("summaries", []),
    }


# ==========================================================
# TRADING212 - VERIFIED REALIZED P/L
# ==========================================================

@router.get("/stocks/trading212/realized")
async def trading212_realized_history():
    result = await safe_collect_trading212_realized_history()
    return {
        "build": BUILD_VERSION,
        "meta": result.get("meta", {}),
        "count": len(result.get("summaries", [])),
        "stocks": result.get("summaries", []),
    }


# ==========================================================
# TRADING212 - HISTORY DEBUG
# ==========================================================

@router.get("/stocks/trading212/history-debug")
async def trading212_history_debug(
    limit: int = 50,
    ticker: str | None = None,
):
    """
    Diagnostic endpoint. Fetches one historical-order page and returns
    representative BUY/SELL rows. Realized P/L is now enabled separately via
    /stocks/trading212/realized and the full update-excel pipeline.
    """
    return await safe_fetch_trading212_history_debug(
        limit=limit,
        ticker=ticker,
    )


# ==========================================================
# CAPITAL - TRANSACTIONS DEBUG
# ==========================================================

@router.get("/stocks/capital/transactions-debug")
async def capital_transactions_debug(
    from_date: str = "2020-01-01T00:00:00",
    to_date: str | None = None,
):
    """
    Diagnostic endpoint only. Returns representative Trade closed and
    commission transactions from Capital so signed P/L and reference linkage
    can be verified before final realized P/L integration.
    """
    return await safe_fetch_capital_transactions_debug(
        from_date=from_date,
        to_date=to_date,
    )


# ==========================================================
# CAPITAL - TEST
# ==========================================================

@router.get("/stocks/capital/test")
async def test_capital():

    result = await create_capital_session()

    return result


# ==========================================================
# CAPITAL - RAW POSITIONS
# ==========================================================

@router.get("/stocks/capital")
async def capital_positions():

    positions = await get_capital_positions()

    return positions


# ==========================================================
# CAPITAL - CLEAN
# ==========================================================

@router.get("/stocks/capital/clean")
async def clean_capital_positions():

    stocks = (
        await get_clean_capital_positions()
    )

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# FREEDOM24 - TEST
# ==========================================================

@router.get("/stocks/freedom/test")
async def test_freedom():

    return test_freedom_connection()


# ==========================================================
# FREEDOM24 - RAW POSITIONS
# ==========================================================

@router.get("/stocks/freedom")
async def freedom_positions():

    positions = get_freedom_positions()

    return {
        "count": len(positions),
        "positions": positions
    }


# ==========================================================
# FREEDOM24 - TRADES
# ==========================================================

@router.get("/stocks/freedom/trades")
async def freedom_trades():

    trades = get_freedom_trades()

    return {
        "count": len(trades),
        "trades": trades
    }


# ==========================================================
# FREEDOM24 - CLEAN
# ==========================================================

@router.get("/stocks/freedom/clean")
async def clean_freedom_positions():

    stocks = get_clean_freedom_positions()

    return {
        "count": len(stocks),
        "stocks": stocks
    }


# ==========================================================
# FREEDOM24 - DEBUG CURRENCIES
# ==========================================================

@router.get("/stocks/freedom/debug")
async def freedom_debug():

    connection = get_freedom_connection()

    response = (
        connection.authorized_request(
            "getPositionJson"
        )
    )

    portfolio = response["result"]["ps"]

    accounts = portfolio.get(
        "acc",
        []
    )

    return {
        "accounts_count":
            len(accounts),

        "account_keys": (
            list(accounts[0].keys())
            if accounts
            else []
        ),

        "currencies": [
            {
                "curr":
                    account.get("curr"),

                "currval":
                    account.get("currval")
            }
            for account in accounts
        ]
    }


# ==========================================================
# FREEDOM24 - TRADES DEBUG
# ==========================================================

@router.get("/stocks/freedom/trades-debug")
async def freedom_trades_debug():

    connection = get_freedom_connection()

    response = (
        connection.authorized_request(
            "getTradesHistory",
            {
                "beginDate":
                    "2020-01-01T00:00:00",

                "endDate":
                    "2030-12-31T23:59:59",

                "max":
                    100,

                "sort":
                    1,
            }
        )
    )

    trades_section = response.get(
        "trades",
        {}
    )

    trades = trades_section.get(
        "trade",
        []
    )

    return {
        "max_trade_id":
            response.get("max_trade_id"),

        "trade_count": (
            len(trades)
            if isinstance(trades, list)
            else 1
        ),

        "first_trade_id": (
            trades[0].get("id")
            if (
                isinstance(trades, list)
                and trades
            )
            else None
        ),

        "last_trade_id": (
            trades[-1].get("id")
            if (
                isinstance(trades, list)
                and trades
            )
            else None
        )
    }
@router.get("/stocks/all")
async def get_all_stocks():

    # Trading212
    trading212_stocks = await get_clean_positions()

    # Capital
    capital_stocks = await get_clean_capital_positions()

    # Freedom24
    freedom_stocks = get_clean_freedom_positions()

    # Όλες οι θέσεις μαζί
    all_stocks = (
        trading212_stocks
        + capital_stocks
        + freedom_stocks
    )
    # Ενώνουμε τις πολλαπλές θέσεις
    # της ίδιας μετοχής στον ίδιο broker
    all_stocks = aggregate_positions(
        all_stocks
    )
    # ==================================================
    # ΣΥΝΟΛΙΚΗ ΑΞΙΑ ΧΑΡΤΟΦΥΛΑΚΙΟΥ
    # Όλα τα market_value είναι πλέον σε EUR
    # ==================================================

    total_portfolio_value = sum(
        stock["market_value"]
        for stock in all_stocks
        if stock.get("market_value") is not None
    )

    # ==================================================
    # PORTFOLIO WEIGHT
    # Βάρος κάθε θέσης στο συνολικό χαρτοφυλάκιο
    # ==================================================

    for stock in all_stocks:

        market_value = stock.get("market_value")

        if (
                market_value is not None
                and total_portfolio_value > 0
        ):

            stock["portfolio_weight"] = round(
                market_value
                / total_portfolio_value
                * 100,
                2
            )

        else:

            stock["portfolio_weight"] = None
    return {
        "count": len(all_stocks),
        "total_portfolio_value": round(
            total_portfolio_value,
            2
        ),
        "stocks": all_stocks
    }


@router.get("/stocks/analysis")
async def portfolio_analysis():

    trading212 = await get_clean_positions()
    capital = await get_clean_capital_positions()
    freedom = get_clean_freedom_positions()

    all_stocks = (
        trading212
        + capital
        + freedom
    )

    aggregated = aggregate_positions(
        all_stocks
    )

    return analyze_portfolio(
        aggregated
    )


@router.get("/stocks/fundamentals/{symbol}")
async def stock_fundamentals(symbol: str):

    return get_stock_fundamentals(
        symbol.upper()
    )

@router.get("/stocks/fundamentals")
def get_portfolio_fundamentals(stocks):

    results = []

    for stock in stocks:

        symbol = stock.get("symbol")
        platform = stock.get("platform")

        if not symbol:
            continue

        # --------------------------------------------------
        # YAHOO SYMBOL
        # --------------------------------------------------

        yahoo_symbol = symbol

        # Freedom24 χρησιμοποιεί δικά της symbols:
        # ACM.US -> ACM
        # ASML.EU -> ASML
        # AEGN.GR -> AEGN.AT
        # κτλ.
        if platform == "Freedom24":

            yahoo_symbol = FREEDOM_TICKER_MAP.get(
                symbol,
                symbol
            )

        try:

            fundamentals = get_stock_fundamentals(
                yahoo_symbol
            )

            results.append({

                # Κρατάμε το αρχικό broker symbol
                "symbol": symbol,

                # Για έλεγχο βλέπουμε και τι στείλαμε Yahoo
                "yahoo_symbol": yahoo_symbol,

                "name": stock.get("name"),
                "sector": stock.get("sector"),
                "country": stock.get("country"),
                "platform": platform,

                "fcf_yield": fundamentals.get(
                    "fcf_yield"
                ),

                "fcf_growth": fundamentals.get(
                    "fcf_growth"
                ),

                "roic": fundamentals.get(
                    "roic"
                ),
            })

        except Exception as e:

            print(
                f"FUNDAMENTALS ERROR "
                f"{symbol} -> {yahoo_symbol}: {e}"
            )

            results.append({

                "symbol": symbol,
                "yahoo_symbol": yahoo_symbol,
                "name": stock.get("name"),
                "sector": stock.get("sector"),
                "country": stock.get("country"),
                "platform": platform,

                "fcf_yield": None,
                "fcf_growth": None,
                "roic": None,
            })

    return results
# ==========================================================
# GLOBAL CANDIDATE RESEARCH -> TOP 20 -> EXCEL Candidates
# ==========================================================

@router.post("/stocks/research-candidates")
async def research_candidate_stocks():
    """
    Candidate research restricted to United States + Europe.

    Final Score:
        internally uses 60% relative + 40% absolute,
        but only the final score is exposed in Excel/dashboard.

    Candidate gates:
        - >=25% below validated 52-week high
        - >=60% absolute-data coverage
        - same-country+sector peers; fallback to global sector
          when fewer than 5 peers
        - one row per company (dual listings deduplicated)
        - Financial Services use dedicated ROE/P-B/growth model
        - secondary 1Y daily-history validation
        - company domicile restricted to United States / Europe
        - news/earnings + independent fundamentals + AI commentary
    """
    try:
        existing_rows = get_excel_stocks()
    except Exception:
        existing_rows = []

    existing_symbols = [
        row.get("symbol")
        for row in existing_rows
        if row.get("symbol")
    ]

    try:
        result = await asyncio.to_thread(
            research_top_candidates,
            limit=20,
            min_discount_pct=25.0,
            existing_symbols=existing_symbols,
            min_absolute_coverage=60.0,
        )
    except Exception as error:
        raise HTTPException(
            status_code=503,
            detail=(
                "Candidate research failed: "
                f"{error}"
            ),
        )

    candidates = result.get(
        "stocks",
        [],
    )

    file_path = (
        update_candidate_research_excel(
            candidates,
            research_meta=result.get(
                "meta"
            ),
        )
    )

    dashboard_files = (
        refresh_dashboard_candidates(
            candidates=candidates
        )
    )

    hidden_fields = {
        "relative_score",
        "absolute_score",
        "absolute_coverage",
        "score_mode",
    }

    public_candidates = []

    for candidate in candidates:
        public_candidates.append({
            key: value
            for key, value
            in candidate.items()
            if key not in hidden_fields
            and not key.startswith(
                "absolute_"
            )
            and not key.startswith(
                "relative_"
            )
            and "_absolute_score" not in key
            and "_relative_score" not in key
        })

    return {
        "message": (
            "US/Europe candidate research V20 completed; "
            "Candidates sheet, peer/event/scenario AI context and dashboard updated"
        ),
        "count": len(
            public_candidates
        ),
        "criteria": {
            "score_display": (
                "Final Score only"
            ),
            "minimum_discount_from_52w_high_pct": 25.0,
            "minimum_absolute_coverage_pct": 60.0,
            "peer_rule": (
                "same country + same sector; "
                "fallback to global same-sector "
                "if fewer than 5 valid peers"
            ),
            "deduplicate_companies": True,
            "excluded_candidate_sectors": [],
            "financial_services_model": (
                "ROE + price-to-book + profit margin + revenue growth + forward EPS growth"
            ),
            "fundamentals_crosscheck": (
                "FMP when configured; SEC EDGAR fallback for US issuers"
            ),
            "news_earnings_check": True,
            "ai_commentary": (
                "OpenAI Responses API when OPENAI_API_KEY exists; deterministic fallback otherwise"
            ),
            "peer_intelligence": (
                "FMP peers + Yahoo metrics; industry/sector fallback"
            ),
            "event_impact": (
                "AI/rules impact -100..+100; separate from Final Score"
            ),
            "scenario_analysis": (
                "EPS -15pp, growth slowdown and bear-case shadow scores"
            ),
            "backtest_learning": (
                "point-in-time snapshots and forward-return IC weight recommendations"
            ),
            "secondary_validation": (
                "1Y daily price history with 52W fallback"
            ),
            "geographic_scope": (
                "United States and Europe only"
            ),
            "dashboard_candidate_analysis": (
                "AI/data commentary with score rationale, news, earnings, anomalies and source cross-check"
            ),
            "history_failure_fallback": (
                "use discovery/detail current_price + high_52w "
                "and re-check the 25% rule"
            ),
        },
        "research": result.get(
            "meta",
            {},
        ),
        "excel": str(
            file_path
        ),
        "dashboard": (
            {
                "json_file": str(
                    dashboard_files[
                        "json_file"
                    ]
                ),
                "html_file": str(
                    dashboard_files[
                        "html_file"
                    ]
                ),
            }
            if dashboard_files
            else None
        ),
        "stocks": public_candidates,
    }


@router.get("/stocks/scores")
async def portfolio_scores():

    # ----------------------------------------------
    # ΠΑΙΡΝΟΥΜΕ ΤΙΣ ΘΕΣΕΙΣ ΑΠΟ ΤΙΣ 3 ΠΛΑΤΦΟΡΜΕΣ
    # ----------------------------------------------

    trading212 = await get_clean_positions()
    capital = await get_clean_capital_positions()
    freedom = get_clean_freedom_positions()

    all_stocks = (
        trading212
        + capital
        + freedom
    )

    # ----------------------------------------------
    # ΟΜΑΔΟΠΟΙΗΣΗ ΘΕΣΕΩΝ
    # ----------------------------------------------

    aggregated = aggregate_positions(
        all_stocks
    )

    aggregated = enrich_stocks_with_saved_metadata(
        aggregated
    )

    # ----------------------------------------------
    # FUNDAMENTALS
    # ----------------------------------------------

    fundamentals = get_portfolio_fundamentals(
        aggregated
    )

    # ----------------------------------------------
    # SCORE ΑΝΑ SECTOR
    # ----------------------------------------------

    scores = score_stocks_by_sector(
        fundamentals
    )

    scores = add_forward_growth_scores(
        scores
    )

    scores = add_absolute_quality_scores(
        scores
    )

    scores = await asyncio.to_thread(
        enrich_with_secondary_fundamentals,
        scores,
    )
    scores = await asyncio.to_thread(
        apply_financial_sector_model,
        scores,
    )

    return {
        "count": len(scores),
        "stocks": scores
    }
