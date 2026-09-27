from datetime import date, datetime
from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pytest
from openpyxl import Workbook, load_workbook

from metoxes.services import excel_service


def make_empty_portfolio(path):
    """Δημιουργεί καθαρό portfolio.xlsx μόνο με headers."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(excel_service.CORE_HEADERS)
    workbook.save(path)
    workbook.close()


def sample_stock(
    symbol="TEST",
    platform="Trading212",
    purchase_date="15/01/2026",
):
    return {
        "symbol": symbol,
        "name": "Test Company",
        "sector": "Technology",
        "country": "USA",
        "platform": platform,
        "purchase_date": purchase_date,
        "purchase_price": 100.0,
        "quantity": 10.0,
        "current_price": 120.0,
        "market_value": 1200.0,
        "percent_change": 20.0,
        "monthly_percent_change": 5.0,
        "portfolio_weight": 100.0,
        "vuaa_return": 4.0,
        "excess_return": 16.0,
    }


def test_sold_sheet_created_and_row_archived(tmp_path, monkeypatch):
    """
    Αν μια παλιά θέση εξαφανιστεί από τα νέα stocks,
    πρέπει να μεταφερθεί στο 2ο φύλλο Sold.
    """
    excel_file = tmp_path / "portfolio.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(excel_service.CORE_HEADERS)

    old_stock = sample_stock(
        symbol="TESTSOLD",
        platform="Trading212",
        purchase_date="01/02/2026",
    )

    sheet.append(
        excel_service.stock_to_excel_row(
            old_stock,
            excel_service.CORE_HEADERS,
        )
    )

    workbook.save(excel_file)
    workbook.close()

    monkeypatch.setattr(
        excel_service,
        "EXCEL_FILE",
        excel_file,
    )

    # Δεν υπάρχει πια καμία ενεργή θέση -> TESTSOLD πουλήθηκε.
    excel_service.update_portfolio_excel([])

    workbook = load_workbook(excel_file)
    try:
        assert workbook.sheetnames[0] == "Portfolio"
        assert workbook.sheetnames[1] == "Sold"

        sold = workbook["Sold"]
        headers = excel_service.get_header_lookup(sold)

        assert sold.cell(
            row=2,
            column=headers["symbol"],
        ).value == "TESTSOLD"

        assert sold.cell(
            row=2,
            column=headers["platform"],
        ).value == "Trading212"

        assert sold.cell(
            row=2,
            column=headers["purchase_price"],
        ).value == 100.0

        assert sold.cell(
            row=2,
            column=headers["quantity"],
        ).value == 10.0
    finally:
        workbook.close()


def test_sold_position_is_not_duplicated(tmp_path, monkeypatch):
    """
    Αν τρέξουμε update δεύτερη φορά,
    η ίδια πωλημένη θέση δεν πρέπει να γραφτεί ξανά στο Sold.
    """
    excel_file = tmp_path / "portfolio.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(excel_service.CORE_HEADERS)

    old_stock = sample_stock(
        symbol="TESTSOLD",
        platform="Capital",
    )

    sheet.append(
        excel_service.stock_to_excel_row(
            old_stock,
            excel_service.CORE_HEADERS,
        )
    )

    workbook.save(excel_file)
    workbook.close()

    monkeypatch.setattr(
        excel_service,
        "EXCEL_FILE",
        excel_file,
    )

    excel_service.update_portfolio_excel([])
    excel_service.update_portfolio_excel([])

    workbook = load_workbook(excel_file)
    try:
        sold = workbook["Sold"]

        symbols = [
            sold.cell(row=r, column=1).value
            for r in range(2, sold.max_row + 1)
            if sold.cell(row=r, column=1).value
        ]

        assert symbols.count("TESTSOLD") == 1
    finally:
        workbook.close()


def test_purchase_date_is_real_excel_date(tmp_path, monkeypatch):
    """
    Η purchase_date πρέπει να αποθηκεύεται σαν πραγματική ημερομηνία
    και όχι σαν text, ώστε το Excel να δίνει Date Filters.
    """
    excel_file = tmp_path / "portfolio.xlsx"
    make_empty_portfolio(excel_file)

    monkeypatch.setattr(
        excel_service,
        "EXCEL_FILE",
        excel_file,
    )

    stock = sample_stock(
        symbol="DATE1",
        purchase_date="26/09/2026",
    )

    excel_service.update_portfolio_excel([stock])

    workbook = load_workbook(excel_file)
    try:
        sheet = workbook["Portfolio"]
        headers = excel_service.get_header_lookup(sheet)

        cell = sheet.cell(
            row=2,
            column=headers["purchase_date"],
        )

        assert isinstance(cell.value, (datetime, date))
        assert cell.value.day == 26
        assert cell.value.month == 9
        assert cell.value.year == 2026
        assert cell.number_format == "dd/mm/yyyy"
    finally:
        workbook.close()


def test_portfolio_table_and_summary_formulas_exist(
    tmp_path,
    monkeypatch,
):
    """
    Μετά το update πρέπει να υπάρχουν:
    - Excel Table StockSummary
    - τα 4 summary κουτάκια με formulas
    """
    excel_file = tmp_path / "portfolio.xlsx"
    make_empty_portfolio(excel_file)

    monkeypatch.setattr(
        excel_service,
        "EXCEL_FILE",
        excel_file,
    )

    stock = sample_stock(symbol="FORMULA")
    excel_service.update_portfolio_excel([stock])

    workbook = load_workbook(
        excel_file,
        data_only=False,
    )

    try:
        sheet = workbook["Portfolio"]

        assert "StockSummary" in sheet.tables

        # 1 data row -> last_data_row = 2
        # title_row = 8, value_row = 9
        assert sheet["B8"].value == "PORTFOLIO VALUE"
        assert sheet["E8"].value == "PORTFOLIO RETURN"
        assert sheet["I8"].value == "VUAA RETURN"
        assert sheet["L8"].value == "EXCESS RETURN"

        assert isinstance(sheet["B9"].value, str)
        assert "StockSummary[current_price]" in sheet["B9"].value

        assert isinstance(sheet["E9"].value, str)
        assert "StockSummary[purchase_price]" in sheet["E9"].value

        assert isinstance(sheet["I9"].value, str)
        assert "StockSummary[vuaa_return]" in sheet["I9"].value

        assert sheet["L9"].value == "=E9-I9"
    finally:
        workbook.close()


def test_avg_monthly_change_example():
    """
    Παράδειγμα χρήστη:
    +30% σε 3 ολοκληρωμένους μήνες = +10% / μήνα.
    """
    result = excel_service.calculate_avg_monthly_change(
        percent_change=30,
        purchase_date="15/01/2026",
        as_of_date=date(2026, 4, 15),
    )

    assert result == 10.0
