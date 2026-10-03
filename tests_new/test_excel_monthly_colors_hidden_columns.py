from openpyxl import Workbook

from metoxes.services.excel_service import (
    CORE_HEADERS,
    OPTIONAL_HEADERS,
    create_summary_boxes,
    format_portfolio_sheet,
)


def test_portfolio_hidden_columns_and_real_monthly_color_bands():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"

    headers = CORE_HEADERS + OPTIONAL_HEADERS
    sheet.append(headers)

    avg_monthly_values = [12, 7, 2, -5, -10]
    symbols = ["A", "B", "C", "D", "E"]

    for symbol, avg_monthly in zip(symbols, avg_monthly_values):
        sheet.append([
            symbol, "Test Company", "Technology", "United States", "Test",
            "2026-01-01", 100, 1, 110, 110, 10, -3, avg_monthly, 1, 5, 5,
            7, 10, 20, 12, 15, 80,
        ])

    format_portfolio_sheet(
        sheet,
        last_data_row=6,
        active_headers=headers,
    )

    lookup = {
        str(cell.value).strip().lower(): cell.column_letter
        for cell in sheet[1]
        if cell.value
    }

    hidden_fields = [
        "sector",
        "purchase_price",
        "quantity",
        "current_price",
        "market_value",
        "fcf_yield",
        "fcf_growth",
        "roic",
        "forward_revenue_growth",
        "forward_eps_growth",
    ]

    for field in hidden_fields:
        assert sheet.column_dimensions[
            lookup[field]
        ].hidden is True

    assert sheet.column_dimensions[
        lookup["monthly_percent_change"]
    ].hidden is not True
    assert sheet.column_dimensions[
        lookup["avg_monthly_change"]
    ].hidden is not True
    assert sheet.column_dimensions[
        lookup["final_score"]
    ].hidden is not True

    expected = [
        "001F4E78",
        "005B9BD5",
        "00D9EAF7",
        "00FCE4D6",
        "00ED7D31",
    ]

    # The entire row must carry the same real fill, driven by column M (avg_monthly_change).
    for row_number, expected_fill in zip(range(2, 7), expected):
        for column in range(1, len(headers) + 1):
            assert (
                sheet.cell(row=row_number, column=column).fill.fgColor.rgb
                == expected_fill
            )

    # No conditional-format rule is needed; the rows have literal fills.
    assert len(list(sheet.conditional_formatting)) == 0

    workbook.close()


def test_vuaa_summary_box_is_visible_in_n_o_columns():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"

    headers = CORE_HEADERS + OPTIONAL_HEADERS
    sheet.append(headers)
    sheet.append([
        "TEST", "Test Company", "Technology", "United States", "Test",
        "2026-01-01", 100, 1, 110, 110, 10, 5, 1, 1, 5, 5,
        7, 10, 20, 12, 15, 80,
    ])

    format_portfolio_sheet(
        sheet,
        last_data_row=2,
        active_headers=headers,
    )
    create_summary_boxes(sheet, 2)

    title_row = 8
    value_row = 9

    assert sheet[f"N{title_row}"].value == "VUAA RETURN"
    assert "StockSummary[vuaa_return]" in sheet[f"N{value_row}"].value
    assert sheet[f"L{value_row}"].value == f"=E{value_row}-N{value_row}"
    assert sheet.column_dimensions["N"].hidden is not True
    assert sheet.column_dimensions["O"].hidden is not True

    workbook.close()


def test_recalculation_handles_missing_calc_properties():
    from metoxes.services.excel_service import set_excel_recalculation

    workbook = Workbook()
    workbook.calculation = None

    set_excel_recalculation(workbook)

    assert workbook.calculation is not None
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True

    workbook.close()


def test_avg_monthly_change_keeps_positive_sign_for_up_positions():
    from datetime import date
    from metoxes.services.excel_service import calculate_avg_monthly_change

    # CRDO-like example: +25.72% total return, position less than one full
    # month old -> minimum denominator 1 -> still positive +25.72%.
    assert calculate_avg_monthly_change(
        25.72, date(2026, 9, 2), date(2026, 10, 1)
    ) == 25.72

    # OptimumBank-like example: +112.85% total return over 12 completed
    # months -> +9.40% average per month, never a negative sign.
    assert calculate_avg_monthly_change(
        112.85, date(2025, 9, 22), date(2026, 10, 1)
    ) == 9.4


def test_row_color_uses_avg_monthly_not_latest_month():
    workbook = Workbook()
    sheet = workbook.active
    headers = CORE_HEADERS + OPTIONAL_HEADERS
    sheet.append(headers)

    # Latest 1M return is negative, but average monthly return is strongly
    # positive. Row must therefore be blue, not orange.
    sheet.append([
        "CRDO", "Credo", "Technology", "Cayman Islands", "Test",
        "2026-09-02", 145.79, 10, 183.28, 1832.8, 25.72, -0.33, 25.72,
        2, 2.5, 23.2, 1.05, 1302.37, 21.42, 54.8, 53.66, 70.6,
    ])

    format_portfolio_sheet(sheet, 2, headers)
    assert sheet["A2"].fill.fgColor.rgb == "001F4E78"
    workbook.close()
