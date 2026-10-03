from openpyxl import Workbook

from metoxes.services.excel_service import (
    compact_sold_sheet,
)


def _make_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    sold = wb.create_sheet("Sold")
    sold.append([
        "symbol", "name", "sector", "country", "platform",
        "purchase_date", "purchase_price", "quantity",
        "current_price", "market_value", "percent_change",
    ])
    sold.append([
        "AEGN.GR", "Aegean", "Industrials", "Greece", "Freedom24",
        "2026-01-01", 12.37, 110, 12.12, 1333.2, -2.02,
    ])
    return wb


def test_compact_sold_uses_verified_broker_pnl():
    wb = _make_workbook()
    lookup = {
        ("AEGN.GR", "FREEDOM24"): {
            "net_profit_loss_eur": -34.37,
            "sale_proceeds_eur": 1333.2,
            "cost_basis_eur": 1361.05,
            "average_sale_price_eur": 12.12,
            "realized_return_pct": -2.53,
        }
    }

    compact_sold_sheet(wb, sold_realized_lookup=lookup)
    sold = wb["Sold"]

    assert sold.max_column == 9
    assert sold["I2"].value == -34.37
    assert sold["G2"].value == 12.12
    assert sold["H2"].value == -2.53

    title_row = 2 + 6
    value_row = title_row + 1
    assert sold.cell(title_row, 2).value == "SOLD VALUE"
    assert sold.cell(value_row, 2).value == 1333.2
    assert sold.cell(title_row, 5).value == "SOLD RETURN"
    assert sold.cell(title_row, 8).value == "REALIZED P/L"
    assert sold.cell(value_row, 8).value == -34.37


def test_compact_sold_does_not_estimate_pnl_without_history():
    wb = _make_workbook()

    compact_sold_sheet(wb, sold_realized_lookup={})
    sold = wb["Sold"]

    assert sold["I2"].value is None

    title_row = 2 + 6
    value_row = title_row + 1
    assert sold.cell(value_row, 2).value == "PENDING VALUE API 0/1"
    assert sold.cell(value_row, 5).value == "PENDING VALUE API 0/1"
    assert sold.cell(value_row, 8).value == "PENDING P/L API 0/1"


def test_capital_pnl_can_show_without_fabricating_sold_value_or_return():
    wb = _make_workbook()
    sold = wb["Sold"]
    sold["A2"] = "IESC"
    sold["C2"] = "Industrials"
    sold["E2"] = "Capital"

    lookup = {
        ("IESC", "CAPITAL"): {
            "net_profit_loss_eur": 360.67,
            "sale_proceeds_eur": None,
            "cost_basis_eur": None,
            "average_sale_price_eur": None,
            "realized_return_pct": None,
        }
    }

    compact_sold_sheet(wb, sold_realized_lookup=lookup)
    sold = wb["Sold"]

    assert sold["I2"].value == 360.67
    title_row = 2 + 6
    value_row = title_row + 1
    assert sold.cell(value_row, 2).value == "PENDING VALUE API 0/1"
    assert sold.cell(value_row, 5).value == "PENDING VALUE API 0/1"
    assert sold.cell(value_row, 8).value == 360.67


def test_update_portfolio_excel_refreshes_existing_sold_sheet(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    from metoxes.services import excel_service

    path = tmp_path / "portfolio.xlsx"
    wb = Workbook()
    portfolio = wb.active
    portfolio.title = "Portfolio"
    portfolio.append(excel_service.CORE_HEADERS)
    portfolio.append([
        "KEEP", "Keep Co", "Technology", "United States", "Trading212",
        "2026-01-01", 10, 1, 11, 11, 10, 1, 1, 100, 5, 5,
    ])
    sold = wb.create_sheet("Sold")
    sold.append([
        "symbol", "name", "sector", "country", "platform",
        "purchase_date", "purchase_price", "quantity", "current_price",
        "market_value", "percent_change",
    ])
    sold.append([
        "IESC", "IES Holdings", "Industrials", "United States", "Capital",
        "2026-09-01", 99.09, 1, 285.7, 285.7, 188.32,
    ])
    wb.save(path)
    wb.close()

    monkeypatch.setattr(excel_service, "EXCEL_FILE", path)

    lookup = {
        ("IESC", "CAPITAL"): {
            "net_profit_loss_eur": 360.67,
            "sale_proceeds_eur": None,
            "cost_basis_eur": None,
        }
    }
    stocks = [{
        "symbol": "KEEP", "name": "Keep Co", "sector": "Technology",
        "country": "United States", "platform": "Trading212",
        "purchase_date": "2026-01-01", "purchase_price": 10,
        "quantity": 1, "current_price": 11, "market_value": 11,
        "percent_change": 10, "monthly_percent_change": 1,
        "avg_monthly_change": 1, "portfolio_weight": 100,
        "vuaa_return": 5, "excess_return": 5,
    }]

    excel_service.update_portfolio_excel(stocks, sold_realized_lookup=lookup)

    wb2 = load_workbook(path, data_only=False)
    sold2 = wb2["Sold"]
    assert sold2.max_column == 9
    assert sold2["A1"].value == "symbol"
    assert sold2["I1"].value == "total_profit_loss"
    assert sold2["I2"].value == 360.67
    wb2.close()
