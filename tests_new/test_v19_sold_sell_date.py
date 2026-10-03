from datetime import datetime

from openpyxl import Workbook

from metoxes.services import excel_service
from metoxes.services.sold_history_service import (
    summarize_freedom_realized_trades,
    summarize_trading212_realized_orders,
)


def _t212_sell(ticker, quantity, price, fill_id, filled_at, pnl=1.0):
    gross = abs(quantity) * price
    return {
        "order": {
            "status": "FILLED",
            "side": "SELL",
            "ticker": ticker,
            "filledQuantity": -abs(quantity),
            "instrument": {"ticker": ticker, "currency": "EUR"},
        },
        "fill": {
            "id": fill_id,
            "type": "TRADE",
            "quantity": -abs(quantity),
            "price": price,
            "filledAt": filled_at,
            "walletImpact": {
                "currency": "EUR",
                "netValue": gross,
                "realisedProfitLoss": pnl,
                "fxRate": 1.0,
                "taxes": [],
            },
        },
    }


def test_trading212_sell_date_uses_day_with_largest_aggregate_sale_value():
    # 01/10 has two partial fills worth 500 + 500 = 1000 EUR.
    # 02/10 has one larger individual fill worth 945 EUR.
    # The requested sell date is therefore 01/10, because fills are grouped by day.
    items = [
        _t212_sell("TESTd_EQ", 5, 100, "f1", "2026-10-01T10:00:00Z"),
        _t212_sell("TESTd_EQ", 5, 100, "f2", "2026-10-01T14:00:00Z"),
        _t212_sell("TESTd_EQ", 9, 105, "f3", "2026-10-02T11:00:00Z"),
    ]

    summary = summarize_trading212_realized_orders(items)[0]

    assert summary["symbol"] == "TEST.DE"
    assert summary["sell_date"] == "2026-10-01"
    assert summary["sell_date_basis"] == "largest_daily_execution_value_eur"
    assert summary["last_sale_date"] == "2026-10-02"


def test_freedom_sell_date_uses_largest_daily_execution_value():
    trades = [
        {
            "type": "2", "instr_type_c": "1", "instr_nm": "ABC",
            "q": 2, "p": 100, "v": 200, "profit": 10,
            "date": "2026-09-10T10:00:00", "curr_c": "EUR",
            "commission": 0, "commission_currency": "EUR",
            "commiss_exchange": 0,
        },
        {
            "type": "2", "instr_type_c": "1", "instr_nm": "ABC",
            "q": 4, "p": 100, "v": 400, "profit": 20,
            "date": "2026-09-11T10:00:00", "curr_c": "EUR",
            "commission": 0, "commission_currency": "EUR",
            "commiss_exchange": 0,
        },
        {
            "type": "2", "instr_type_c": "1", "instr_nm": "ABC",
            "q": 1, "p": 100, "v": 100, "profit": 5,
            "date": "2026-09-12T10:00:00", "curr_c": "EUR",
            "commission": 0, "commission_currency": "EUR",
            "commiss_exchange": 0,
        },
    ]

    summary = summarize_freedom_realized_trades(trades)[0]
    assert summary["sell_date"] == "2026-09-11"
    assert summary["last_sale_date"] == "2026-09-12"


def test_sold_sheet_has_sell_date_column_and_keeps_quantity_visible():
    wb = Workbook()
    portfolio = wb.active
    portfolio.title = "Portfolio"
    portfolio.append(list(excel_service.CORE_HEADERS))

    current = [{
        "symbol": "CRDO",
        "sector": "Technology",
        "platform": "Trading212",
        "purchase_date": datetime(2026, 9, 2),
        "purchase_price": 145.8,
        "quantity": 0.6,
        "current_price": 194.47,
    }]
    lookup = {
        ("CRDO", "TRADING212"): {
            "symbol": "CRDO",
            "platform": "Trading212",
            "sold_quantity": 9.4,
            "average_execution_price_eur": 193.74,
            "net_profit_loss_eur": 450.75,
            "sell_date": "2026-09-28",
            "last_sale_date": "2026-10-02",
        }
    }

    excel_service.compact_sold_sheet(
        wb,
        current_stocks=current,
        sold_realized_lookup=lookup,
    )

    sold = wb["Sold"]
    headers = [
        sold.cell(row=1, column=c).value
        for c in range(1, len(excel_service.SOLD_VIEW_HEADERS) + 1)
    ]

    assert headers == [
        "symbol", "sector", "platform", "purchase_date", "sell_date",
        "purchase_price", "quantity", "current_price", "percent_change",
        "total_profit_loss",
    ]
    assert isinstance(sold["E2"].value, datetime)
    assert sold["E2"].value.date().isoformat() == "2026-09-28"
    assert sold["E2"].number_format == "dd/mm/yyyy"
    assert sold["G2"].value == 9.4
    assert sold.column_dimensions["G"].hidden is False
    assert sold.tables[excel_service.SOLD_TABLE_NAME].ref == "A1:J2"
