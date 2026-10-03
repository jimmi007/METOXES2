from openpyxl import Workbook, load_workbook

import metoxes.services.excel_service as excel_service


def test_candidates_sheet_is_created_with_final_score_only(tmp_path):
    excel_file = tmp_path / "portfolio.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(
        excel_service.CORE_HEADERS
    )
    workbook.save(excel_file)
    workbook.close()

    excel_service.EXCEL_FILE = excel_file

    candidates = [
        {
            "rank": 1,
            "symbol": "TEST",
            "name": "Test Company",
            "sector": "Technology",
            "country": "United States",
            "currency": "USD",
            "current_price": 70,
            "high_52w": 100,
            "price_25pct_below_high": 75,
            "drop_from_52w_high_amount": 30,
            "discount_from_52w_high_pct": 30,
            "meets_25pct_discount": True,
            "already_in_portfolio": False,
            "fcf_yield": 8,
            "fcf_growth": 12,
            "roic": 20,
            "forward_revenue_growth": 10,
            "forward_eps_growth": 15,
            "relative_score": 80,
            "absolute_score": 70,
            "absolute_coverage": 100,
            "score_mode": "relative_60_absolute_40",
            "final_score": 76,
            "market_cap": 5_000_000_000,
            "validation_status": "validated",
            "validation_source": "Yahoo Finance 1Y daily history",
            "researched_at": "2026-09-30T17:00:00",
            "source_url": "https://finance.yahoo.com/quote/TEST",
            "dashboard_analysis": "Σύντομη ανάλυση υποψήφιας μετοχής για το dashboard.",
        }
    ]

    excel_service.update_candidate_research_excel(
        candidates
    )

    workbook = load_workbook(
        excel_file,
        data_only=False,
    )

    sheet = workbook["Candidates"]
    headers = [
        cell.value
        for cell in sheet[1]
    ]

    assert "current_price" in headers
    assert "high_52w" in headers
    assert "price_25pct_below_high" in headers
    assert "discount_from_52w_high_pct" in headers
    assert "final_score" in headers
    assert "dashboard_analysis" in headers

    assert "relative_score" not in headers
    assert "absolute_score" not in headers
    assert "absolute_coverage" not in headers
    assert "score_mode" not in headers

    symbol_col = (
        headers.index("symbol") + 1
    )
    final_col = (
        headers.index("final_score") + 1
    )

    assert (
        sheet.cell(
            2,
            symbol_col,
        ).value
        == "TEST"
    )
    assert (
        sheet.cell(
            2,
            final_col,
        ).value
        == 76
    )

    analysis_col = headers.index(
        "dashboard_analysis"
    ) + 1
    analysis_letter = sheet.cell(
        1, analysis_col
    ).column_letter
    assert sheet.column_dimensions[
        analysis_letter
    ].hidden is True

    visible_letters = {"A", "B", "C", "D", "E", "G", "S", "X"}
    for column_number in range(1, len(headers) + 1):
        letter = sheet.cell(1, column_number).column_letter
        expected_hidden = letter not in visible_letters
        assert sheet.column_dimensions[letter].hidden is expected_hidden

    workbook.close()
