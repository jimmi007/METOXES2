from openpyxl import Workbook

from metoxes.services import excel_service


def _sample_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    headers = list(excel_service.CORE_HEADERS)
    sheet.append(headers)
    sheet.append([
        "AAA", "AAA Corp", "Technology", "USA", "Trading212",
        "01/01/2026", 100.0, 10.0, 120.0, 1200.0,
        20.0, 5.0, 2.0, 100.0, 4.0, 16.0,
    ])
    return workbook, sheet, headers


def test_v18_quantity_is_visible_after_formatting():
    workbook, sheet, headers = _sample_sheet()
    try:
        # Simulate the V17 workbook where quantity had already been hidden.
        quantity_col = headers.index("quantity") + 1
        from openpyxl.utils import get_column_letter
        quantity_letter = get_column_letter(quantity_col)
        sheet.column_dimensions[quantity_letter].hidden = True

        excel_service.format_portfolio_sheet(sheet, 2, headers)
        assert sheet.column_dimensions[quantity_letter].hidden is False
    finally:
        workbook.close()


def test_v18_total_boxes_use_realized_pnl_and_existing_cost():
    workbook, sheet, headers = _sample_sheet()
    try:
        excel_service.format_portfolio_sheet(sheet, 2, headers)
        excel_service.create_summary_boxes(sheet, 2)

        # last_data_row=2 -> title/value 8/9, then three rows lower -> 12/13.
        assert sheet["B12"].value == "TOTAL VALUE"
        assert sheet["E12"].value == "TOTAL PORTFOLIO RETURN"
        assert "SUM(SoldSummary[total_profit_loss])" in sheet["B13"].value
        assert "StockSummary[purchase_price]" in sheet["E13"].value
        assert "StockSummary[quantity]" in sheet["E13"].value
        assert sheet["B13"].number_format == '#,##0.00 "€"'
        assert sheet["E13"].number_format == "0.00%"
    finally:
        workbook.close()
